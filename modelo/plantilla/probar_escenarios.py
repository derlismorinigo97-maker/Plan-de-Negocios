# -*- coding: utf-8 -*-
"""
Batería de escenarios sobre la plantilla v2 (Ecostar).

Cada escenario modifica solo entradas VIGENTES (como lo haría el usuario); el presupuesto original queda como referencia.
Recalcula con LibreOffice, extrae indicadores y ejecuta verificaciones de consistencia del motor.

Uso: python3 probar_escenarios.py libro_sin_recalcular.xlsx recalc.py carpeta_salida
"""
import os
import pickle
import subprocess
import sys
import json
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

BASE, RECALC, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(OUT, exist_ok=True)
FC, N, FIN = 6, 144, 120
LP = dict(monto=2791587, tasa=0.09, gracia=6, cuotas=54)
TC_BASE = 7300
TASA_CP = 0.10        # ilustrativa: entre el promedio activo ME del BCP (7,99%, mar-2026) y el tope (11,24%, jul-2026)


# ---------------------------------------------------------------- acceso por nombre
def cells(wb, name):
    sh, rg = wb.defined_names[name].attr_text.split("!")
    ws = wb[sh.strip("'")]
    rg = rg.replace("$", "")
    a, b = (rg.split(":") + [rg])[:2]
    ca, ra = coordinate_from_string(a)
    cb, rb = coordinate_from_string(b)
    return [(ws, r, c) for r in range(ra, rb + 1) for c in range(column_index_from_string(ca), column_index_from_string(cb) + 1)]


def getv(wb, name):
    return [ws.cell(r, c).value for ws, r, c in cells(wb, name)]


def setv(wb, name, val):
    cs = cells(wb, name)
    vals = val if isinstance(val, list) else [val] * len(cs)
    for (ws, r, c), v in zip(cs, vals):
        ws.cell(r, c).value = v


def scale(src, dst, k):
    def f(wb):
        vals = getv(wb, src)
        cs = cells(wb, dst)
        for (ws, r, c), v in zip(cs, vals):
            if isinstance(v, (int, float)):
                ws.cell(r, c).value = v * k
    return f


def S(name, val):
    return lambda wb: setv(wb, name, val)


def loan2(monto, mes, tasa, gracia, cuotas):
    return [S("V_L2_Monto", monto), S("V_L2_Mes", mes), S("V_L2_Tasa", tasa), S("V_L2_Gracia", gracia), S("V_L2_Cuotas", cuotas)]


def linea(lim, tasa=TASA_CP, desde=1, hasta=None):
    return [S("V_LC_Lim", lim), S("V_LC_Tasa", tasa), S("V_LC_Desde", desde), S("V_LC_Hasta", hasta)]


def capex_extra(k):
    def f(wb):
        for ws, r, c in cells(wb, "Inv_D"):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                ws.cell(r, 10).value = v * k          # «Por contratar (si difiere)»
    return f


def rampa(wb):
    vol = getv(wb, "Vol_O")
    nuevo = list(vol)
    # filas de 10 años: gabinetes (0..9) y cajas (10..19)
    for i, (y1, y2, y3) in enumerate([(15000, 15000, 21000), (90000, 90000, 120000)]):
        nuevo[i * 10 + 0], nuevo[i * 10 + 1], nuevo[i * 10 + 2] = y1, y2, y3
    for (ws, r, c), v, o in zip(cells(wb, "Vol_V"), nuevo, vol):
        if isinstance(o, (int, float)):
            ws.cell(r, c).value = v


MAS, MENOS = 3800000, 1800000
VAR = lambda k: [scale("MP_O", "MP_V", k), scale("INS_O", "INS_V", k), scale("OV_O", "OV_V", k)]

# ---------------------------------------------------------------- escenarios
# id, grupo, nombre, cambios (texto), edits, tipo de verificación, meta
ESC = [
    ("E00", "Referencia", "Plan base (Oct-2025)", "Sin cambios: vigente = presupuesto.", [], "base", {}),
    # ------------------------------------------------ financiamiento x cobranza
    ("F1", "Financiamiento y cobranza", "Más financiamiento", "Préstamo LP USD 3,8 M (+1,0 M), mismas condiciones.", [S("V_L1_Monto", MAS)], "fin", {"lp": MAS}),
    ("F2", "Financiamiento y cobranza", "Menos financiamiento", "Préstamo LP USD 1,8 M (−1,0 M).", [S("V_L1_Monto", MENOS)], "fin", {"lp": MENOS}),
    ("F3", "Financiamiento y cobranza", "Cobro a 180 días", "Días de cobro 120 → 180.", [S("vg_DSO", 180)], "dso", {"dso": 180}),
    ("F4", "Financiamiento y cobranza", "Cobro al contado", "Días de cobro 120 → 0.", [S("vg_DSO", 0)], "dso", {"dso": 0}),
    ("F5", "Financiamiento y cobranza", "Más financiamiento + cobro 180 días", "F1 + F3.", [S("V_L1_Monto", MAS), S("vg_DSO", 180)], "otro", {"lp": MAS}),
    ("F6", "Financiamiento y cobranza", "Menos financiamiento + contado", "F2 + F4.", [S("V_L1_Monto", MENOS), S("vg_DSO", 0)], "otro", {"lp": MENOS}),
    ("F7", "Financiamiento y cobranza", "Más financiamiento + contado", "F1 + F4.", [S("V_L1_Monto", MAS), S("vg_DSO", 0)], "otro", {"lp": MAS}),
    ("F8", "Financiamiento y cobranza", "Menos financiamiento + cobro 180 días", "F2 + F3.", [S("V_L1_Monto", MENOS), S("vg_DSO", 180)], "otro", {"lp": MENOS}),
    # ------------------------------------------------ costos
    ("C1", "Costos", "Costos variables +10%", "Materia prima, insumos y otros variables por unidad +10%.", VAR(1.10), "otro", {}),
    ("C2", "Costos", "Costos variables +20%", "Idem +20%.", VAR(1.20), "otro", {}),
    ("C3", "Costos", "Costos variables +30%", "Idem +30%.", VAR(1.30), "otro", {}),
    ("C4", "Costos", "Costos que no se cubren con las ventas",
     "Variables +30%, salarios +30% y costos fijos ×2.", VAR(1.30) + [scale("Sal_O", "Sal_V", 1.30), scale("FC_O", "FC_V", 2.0)], "otro", {}),
    # ------------------------------------------------ ventas
    ("V1", "Ventas", "Volumen −10%", "Unidades vendidas −10% todos los años.", [scale("Vol_O", "Vol_V", 0.90)], "otro", {}),
    ("V2", "Ventas", "Volumen −20%", "Idem −20%.", [scale("Vol_O", "Vol_V", 0.80)], "otro", {}),
    ("V3", "Ventas", "Volumen −30%", "Idem −30%.", [scale("Vol_O", "Vol_V", 0.70)], "otro", {}),
    ("V4", "Ventas", "Vendo más: volumen +10%", "Unidades +10% (desde el año 4 supera la capacidad instalada).", [scale("Vol_O", "Vol_V", 1.10)], "otro", {}),
    ("V5", "Ventas", "Vendo más dentro de la capacidad", "Años 1-2 a capacidad inicial (15.000 / 90.000), año 3 = 21.000 / 120.000.", [rampa], "otro", {}),
    # ------------------------------------------------ plan no sale como se esperaba
    ("P1", "Plan no sale como se esperaba", "Desfavorable combinado",
     "Volumen −30%, precio −5%, inicio 6 meses después, CAPEX +15%.",
     [scale("Vol_O", "Vol_V", 0.70), scale("Precio_O", "Precio_V", 0.95), S("vg_Inicio", 7), capex_extra(1.15)], "otro", {}),
    ("P2", "Plan no sale como se esperaba", "Selector «Conservador» del modelo",
     "Volumen −10%, precio −5%, materiales +10%, atraso 3 meses, CAPEX pendiente +10%, cobro 150 días.", [S("n_Escenario", "Conservador")], "otro", {}),
    ("P3", "Plan no sale como se esperaba", "Selector «Optimista» del modelo",
     "Volumen +5%, precio +2%, materiales −3%, cobro e inventario 90 días.", [S("n_Escenario", "Optimista")], "otro", {}),
    # ------------------------------------------------ plazo del préstamo
    ("L1", "Plazo del préstamo", "LP a 3 años (6 meses de gracia + 30 cuotas)", "Mismo monto y tasa.", [S("V_L1_Cuotas", 30)], "fin", {"lp": LP["monto"]}),
    ("L2", "Plazo del préstamo", "LP a 7 años (6 + 78 cuotas)", "Mismo monto y tasa.", [S("V_L1_Cuotas", 78)], "fin", {"lp": LP["monto"]}),
    ("L3", "Plazo del préstamo", "LP a 8 años (12 + 84 cuotas)", "Mismo monto y tasa.", [S("V_L1_Gracia", 12), S("V_L1_Cuotas", 84)], "fin", {"lp": LP["monto"]}),
    ("L4", "Plazo del préstamo", "Repago acelerado con todo el excedente",
     "Sin cuotas fijas: línea de USD 4 M al 9% que se usa según la necesidad y se devuelve con todo el excedente de caja.",
     [S("V_L1_Monto", 0)] + linea(4000000, 0.09), "fin", {"lim": 4000000}),
    # ------------------------------------------------ corto vs largo plazo
    ("K1", "Corto vs largo plazo", "Solo LP, ampliado para cubrir la brecha", "Préstamo LP USD 3,25 M (6 + 54) al 9%.", [S("V_L1_Monto", 3250000)], "fin", {"lp": 3250000}),
    ("K2", "Corto vs largo plazo", "LP + préstamo CP de 12 meses",
     "LP base + CP USD 500 mil al 10% en el mes 25, 12 cuotas.", loan2(500000, 25, TASA_CP, 0, 12), "fin", {"lp": LP["monto"] + 500000}),
    ("K3", "Corto vs largo plazo", "LP + línea rotativa a la par", "LP base + línea USD 1 M al 10%, renovada hasta el final.", linea(1000000), "fin", {"lim": 1000000}),
    ("K4", "Corto vs largo plazo", "Solo CP: línea rotativa renovada", "Sin LP; línea USD 4 M al 10% renovada hasta el final.", [S("V_L1_Monto", 0)] + linea(4000000), "fin",
     {"lim": 4000000}),
    ("K5", "Corto vs largo plazo", "Solo CP sin renovación", "Sin LP; línea USD 4 M al 10% que vence en el mes 24.", [S("V_L1_Monto", 0)] + linea(4000000, hasta=24), "fin",
     {"lim": 4000000}),
    ("K6", "Corto vs largo plazo", "LP para CAPEX + CP para capital de trabajo",
     "LP USD 0,76 M (6 + 54) al 9% + línea USD 3 M al 10% renovada.", [S("V_L1_Monto", 760000)] + linea(3000000), "fin", {"lim": 3000000}),
    ("K7", "Corto vs largo plazo", "Solo CP renovado con tasa 7,99%", "K4 con la tasa activa promedio ME del BCP (mar-2026).", [S("V_L1_Monto", 0)] + linea(4000000, 0.0799),
     "fin", {"lim": 4000000}),
    ("K8", "Corto vs largo plazo", "Solo CP renovado con tasa 11,24%", "K4 con el tope legal ME del BCP (jul-2026).", [S("V_L1_Monto", 0)] + linea(4000000, 0.1124), "fin",
     {"lim": 4000000}),
    # ------------------------------------------------ tipo de cambio
    ("T1", "Tipo de cambio", "Sube el TC: Gs 8.760 (+20%)", "TC vigente de todos los años.", [S("TC_V", 8760)], "tc", {"tc": 8760}),
    ("T2", "Tipo de cambio", "Baja el TC: Gs 6.145 (−16%)", "Nivel usado en el archivo de Alianza del Acero (jun-2026).", [S("TC_V", 6145)], "tc", {"tc": 6145}),
    ("T3", "Tipo de cambio", "Baja el TC: Gs 5.840 (−20%)", "TC vigente de todos los años.", [S("TC_V", 5840)], "tc", {"tc": 5840}),
]


# ---------------------------------------------------------------- ejecución
def run(e):
    eid, grupo, nombre, cambios, edits, tipo, meta = e
    wb = load_workbook(BASE)
    for f in edits:
        f(wb)
    path = os.path.join(OUT, f"{eid}.xlsx")
    wb.save(path)
    res = subprocess.run(["python3", RECALC, path, "600"], capture_output=True, text=True)
    j = json.loads(res.stdout[res.stdout.index("{"):])
    wv = load_workbook(path, data_only=True)
    ws = wv["Calculo"]
    rows = {}
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if isinstance(k, str) and "." in k:
            rows[k] = [ws.cell(r, FC + m - 1).value for m in range(1, N + 1)]
    rs = wv["Resultados"]
    kpi = {}
    for r in range(6, 40):
        lab = rs.cell(r, 2).value
        if lab and rs.cell(r, 4).value is not None:
            kpi[lab] = (rs.cell(r, 3).value, rs.cell(r, 4).value)
    ctl = []
    wc = wv["Control"]
    for r in range(5, 30):
        if wc.cell(r, 5).value in ("OK", "REVISAR"):
            ctl.append((wc.cell(r, 1).value, wc.cell(r, 2).value, wc.cell(r, 3).value, wc.cell(r, 5).value))
    os.remove(path)
    return eid, dict(grupo=grupo, nombre=nombre, cambios=cambios, tipo=tipo, meta=meta, err=j.get("total_errors"), rows=rows, kpi=kpi, ctl=ctl)


if __name__ == "__main__":
    with ThreadPoolExecutor(int(os.environ.get("HILOS", "6"))) as ex:
        R = dict(ex.map(run, ESC))
    R = {e[0]: R[e[0]] for e in ESC}
    pickle.dump(R, open(os.path.join(OUT, "escenarios.pkl"), "wb"))
    for k, v in R.items():
        print(k, v["nombre"], "errores:", v["err"])
