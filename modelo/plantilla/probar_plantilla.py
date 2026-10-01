"""Pruebas de la plantilla v2: recalcula casos con LibreOffice y guarda los resultados en tests.pkl.
Uso: python3 probar_plantilla.py libro.xlsx recalc.py carpeta_salida"""
import sys, json, subprocess, os, re
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as CL
BASE, RECALC, SP = sys.argv[1], sys.argv[2], sys.argv[3]   # libro sin recalcular, script de recálculo LibreOffice, carpeta de salida
os.makedirs(SP, exist_ok=True)
FC=6
def ref(wb,name):
    t=wb.defined_names[name].attr_text
    sh,rg=t.split("!"); sh=sh.strip("'"); rg=rg.replace("$","")
    return sh,rg
def setn(wb,name,val,idx=None):
    sh,rg=ref(wb,name); ws=wb[sh]
    if ":" not in rg: ws[rg]=val; return
    a,b=rg.split(":")
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    ca,ra=coordinate_from_string(a); cb,rb=coordinate_from_string(b)
    ca=column_index_from_string(ca); cb=column_index_from_string(cb)
    cells=[(r,c) for r in range(ra,rb+1) for c in range(ca,cb+1)]
    if idx is None:
        for r,c in cells: ws.cell(r,c,val)
    else:
        r,c=cells[idx]; ws.cell(r,c,val)
def rrow(ws,code):
    for r in range(1,ws.max_row+1):
        if ws.cell(r,1).value==code: return r
    raise KeyError(code)
def setreal(wb,code,m,val):
    ws=wb["Reales"]; ws.cell(rrow(ws,code),FC+m-1,val)
REALES={  # mes 1..6
 "TC":[7400]*6,"U1":[600]*6,"U2":[4000]*6,"S1":[162000]*6,"S2":[92000]*6,"MAT":[120000]*6,"OVAR":[30000]*6,
 "PERP":[300000000]*6,"PERA":[60000000]*6,"FIJP":[15000]*6,"ADM":[25000]*6,"TRIB":[2540]*6,
 "CAPEX":[600000,0,0,30000,0,0],"DESEMB":[2791587,0,0,0,0,0],"INT":[23026]*6,"AMORT":[0]*6,"APORTE":[767782,0,0,0,0,0],"DIV":[0]*6}
SALDOS={"S_CAJA":1500000,"S_CXC":900000,"S_INV":500000,"S_IVA":0,"S_PROV":100000,"S_DEUDA":2791587}
def with_reales(wb):
    setn(wb,"n_Corte",6)
    for k,v in REALES.items():
        for m,x in enumerate(v,1): setreal(wb,k,m,x)
    for k,v in SALDOS.items(): setreal(wb,k,6,v)
    # inversión por ítem: pagado = 630000 repartido en primeros ítems
    ws=wb["Inversion"]; r0=rrow(ws,1)
    for i in range(11): ws.cell(r0+i,7).value=ws.cell(r0+i,4).value; ws.cell(r0+i,8).value=ws.cell(r0+i,4).value
    ws.cell(r0+11,7).value=35833.91; ws.cell(r0+11,8).value=35833.91
CASES={
 "base":[],
 "reales_vig":[("reales",)],
 "reales_cons":[("reales",),("n","n_Escenario","Conservador")],
 "tc":[("n","TC_V",8030)],
 "dso":[("n","vg_DSO",150)],
 "cons":[("n","n_Escenario","Conservador")],
 "impl":[("n","o_Inicio",7)],
}
def build(name,edits):
    wb=load_workbook(BASE)
    for e in edits:
        if e[0]=="reales": with_reales(wb)
        else: setn(wb,e[1],e[2])
    out=os.path.join(SP,f"t_{name}.xlsx"); wb.save(out)
    res=subprocess.run(["python3",RECALC,out,"600"],capture_output=True,text=True)
    j=json.loads(res.stdout[res.stdout.index("{"):])
    return name,out,j
with ThreadPoolExecutor(4) as ex:
    outs=list(ex.map(lambda kv: build(*kv), CASES.items()))
R={}
for name,out,j in outs:
    wb=load_workbook(out,data_only=True); ws=wb["Calculo"]
    rows={}
    for r in range(1,ws.max_row+1):
        k=ws.cell(r,1).value
        if isinstance(k,str) and "." in k: rows[k]=[ws.cell(r,FC+m-1).value for m in range(1,145)]
    k=wb["Resultados"]; kp={}
    for r in range(6,40):
        if k.cell(r,2).value: kp[k.cell(r,2).value]=(k.cell(r,3).value,k.cell(r,4).value)
    ctl=[(wb["Control"].cell(r,2).value,wb["Control"].cell(r,3).value,wb["Control"].cell(r,5).value) for r in range(5,22) if wb["Control"].cell(r,5).value]
    R[name]=dict(err=j.get("total_errors"),rows=rows,kpi=kp,ctl=ctl)
import pickle; pickle.dump(R,open(os.path.join(SP,"tests.pkl"),"wb"))
for n in R: print(n,"errores:",R[n]["err"])
