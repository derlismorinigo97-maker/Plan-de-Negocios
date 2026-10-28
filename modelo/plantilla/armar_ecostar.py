# -*- coding: utf-8 -*-
"""
Arma el libro de Ecostar: genera el modelo con generar_plantilla.py y le agrega las hojas de registro
(«Claude Log» al principio y «Esp_25Productos» al final) copiadas de un libro anterior.
Opcional: agrega una fila al «Claude Log» desde un archivo JSON con las claves
turno, fecha (dd/mm/aaaa), solicitud, accion, detalle, resultado.

Uso: python3 armar_ecostar.py salida.xlsx libro_con_registro.xlsx [fila_log.json]
Después hay que recalcular (Excel o LibreOffice headless).
"""
import sys, json, datetime
from copy import copy
from openpyxl import load_workbook

OUT, FUENTE = sys.argv[1], sys.argv[2]
FILA = json.load(open(sys.argv[3], encoding="utf-8")) if len(sys.argv) > 3 else None
sys.argv = [sys.argv[0], "ecostar", OUT]
import generar_plantilla as g   # noqa: E402  (lee el modo y la salida de sys.argv)


def copiar(src, dst):
    for row in src.iter_rows():
        for c in row:
            d = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                d.font, d.fill, d.border = copy(c.font), copy(c.fill), copy(c.border)
                d.alignment, d.number_format, d.protection = copy(c.alignment), c.number_format, copy(c.protection)
    for k, dim in src.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width
    for k, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[k].height = dim.height
    for m in src.merged_cells.ranges:
        dst.merge_cells(str(m))
    dst.freeze_panes = src.freeze_panes
    dst.sheet_properties.tabColor = src.sheet_properties.tabColor
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines


def agregar_hojas(wb):
    fu = load_workbook(FUENTE)
    log = wb.create_sheet("Claude Log", 0)
    copiar(fu["Claude Log"], log)
    if FILA:
        r = log.max_row + 1
        prev = log.cell(row=r - 1, column=1)
        vals = [FILA["turno"], datetime.datetime.strptime(FILA["fecha"], "%d/%m/%Y"), FILA["solicitud"], FILA["accion"], FILA["detalle"], FILA["resultado"]]
        for c, v in enumerate(vals, 1):
            d = log.cell(row=r, column=c, value=v)
            s = log.cell(row=r - 1, column=c)
            if s.has_style:
                d.font, d.fill, d.border, d.alignment, d.number_format = copy(s.font), copy(s.fill), copy(s.border), copy(s.alignment), s.number_format
        if log.row_dimensions[prev.row].height:
            log.row_dimensions[r].height = log.row_dimensions[prev.row].height
    esp = wb.create_sheet("Esp_25Productos")
    copiar(fu["Esp_25Productos"], esp)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    wb.active = wb.sheetnames.index("Inicio")
    wb["Inicio"].sheet_view.tabSelected = True


_save = g.wb.save
g.wb.save = lambda path: (agregar_hojas(g.wb), _save(path))
g.main()
