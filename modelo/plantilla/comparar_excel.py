# -*- coding: utf-8 -*-
"""
Compara un libro validado y guardado en Excel con el libro de referencia recalculado en LibreOffice.
Uso: python3 comparar_excel.py referencia.xlsx libro_excel.xlsx [escenarios.pkl]

1. Fórmulas y constantes de todas las hojas del modelo: detecta restauraciones incompletas.
2. Nombres definidos.
3. Valores calculados por Excel frente a LibreOffice (caso base, celda por celda).
4. Si se pasa escenarios.pkl y el libro tiene la hoja «Informe_Cierre_v6»: los indicadores de cada escenario
   informados por Claude en Excel frente a los de LibreOffice.
"""
import sys, pickle, collections
from openpyxl import load_workbook
REF, XL = sys.argv[1:3]
PKL = sys.argv[3] if len(sys.argv) > 3 else None
NO_MODELO = {"Claude Log", "Esp_25Productos", "Informe_Cierre_v6"}
TOL = 1e-6

def iguales(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= TOL * max(1, abs(a))
    return a == b or (a in (None, "") and b in (None, ""))

def comparar(wa, wb, hojas):
    dif = collections.Counter(); ej = collections.defaultdict(list); n = 0
    for h in hojas:
        a, b = wa[h], wb[h]
        for r in range(1, max(a.max_row, b.max_row) + 1):
            for c in range(1, max(a.max_column, b.max_column) + 1):
                va, vb = a.cell(r, c).value, b.cell(r, c).value
                if va is None and vb is None:
                    continue
                n += 1
                if not iguales(va, vb):
                    dif[h] += 1
                    if len(ej[h]) < 5:
                        ej[h].append((a.cell(r, c).coordinate, va, vb))
    return n, dif, ej

def informe(titulo, n, dif, ej):
    print(f"{titulo}: {n} celdas, {sum(dif.values())} distintas")
    for h in dif:
        print(f"  {h}: {dif[h]}", ej[h])

fa, fb = load_workbook(REF), load_workbook(XL)
hojas = [h for h in fa.sheetnames if h not in NO_MODELO]
faltan = [h for h in hojas if h not in fb.sheetnames]
if faltan:
    sys.exit(f"Faltan hojas en el libro de Excel: {faltan}")
informe("1. Fórmulas y constantes", *comparar(fa, fb, hojas))
na = {k: v.attr_text for k, v in fa.defined_names.items()}
nb = {k: v.attr_text for k, v in fb.defined_names.items()}
print(f"2. Nombres definidos: {len(na)} / {len(nb)}, distintos: {sorted(k for k in na.keys() | nb.keys() if na.get(k) != nb.get(k))}")
va, vb = load_workbook(REF, data_only=True), load_workbook(XL, data_only=True)
informe("3. Valores calculados (Excel frente a LibreOffice)", *comparar(va, vb, hojas))

if PKL and "Informe_Cierre_v6" in vb.sheetnames:
    R = pickle.load(open(PKL, "rb"))
    ws = vb["Informe_Cierre_v6"]
    fila = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Id" and ws.cell(r, 2).value == "Ventas")
    kpis = ["Ventas totales del horizonte", "VAN del proyecto (FCFF a WACC)", "VAN de los socios (a Ke)",
            "Aporte adicional necesario para no tener caja negativa", "Meses con brecha",
            "Intereses, IVA y comisiones totales", "Cancelación total de la deuda (años desde el mes 1)",
            "DSCR mínimo desde el 2.º año de operación (flujo operativo / servicio de deuda)", "Balance de control (0 = OK)"]
    n = 0; malos = []; esc = []
    for r in range(fila + 1, ws.max_row + 1):
        i = ws.cell(r, 1).value
        if i not in R or not isinstance(ws.cell(r, 2).value, (int, float)):
            continue
        esc.append(i); k = dict(R[i]["kpi"])
        for j, nombre in enumerate(kpis):
            n += 1; x, y = ws.cell(r, 2 + j).value, k[nombre][1]
            if not iguales(y, x):
                malos.append((i, nombre, x, y))
    print(f"4. Escenarios informados por Claude en Excel: {len(esc)} escenarios, {n} valores, {len(malos)} distintos")
    for m in malos[:20]:
        print("  ", m)
