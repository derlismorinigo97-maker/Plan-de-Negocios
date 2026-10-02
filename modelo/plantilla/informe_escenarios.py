# -*- coding: utf-8 -*-
"""
Arma el libro de resultados de la batería de escenarios (valores fijos, no fórmulas) y ejecuta las verificaciones.
Uso: python3 informe_escenarios.py escenarios.pkl salida.xlsx [referencia.pkl]
     referencia.pkl (opcional): corrida anterior para la prueba de regresión escenario por escenario.
     Variables de entorno opcionales: MODELO (título), LIBRO (archivo probado), FECHA (fecha de la corrida).
"""
import os
import pickle
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as CL

R = pickle.load(open(sys.argv[1], "rb"))
OUT = sys.argv[2]
REF = pickle.load(open(sys.argv[3], "rb")) if len(sys.argv) > 3 else None
MODELO = os.environ.get("MODELO", "Ecostar v2")
LIBRO = os.environ.get("LIBRO", "Modelo_Maquila_ECOSTAR_v2.xlsx")
FECHA = os.environ.get("FECHA", "01/10/2026")
FIN, N = 120, 144
B = "E00"

F = Font(name="Arial", size=9)
FB = Font(name="Arial", size=9, bold=True)
FH = Font(name="Arial", size=9, bold=True, color="FFFFFF")
FT = Font(name="Arial", size=13, bold=True, color="1F3864")
FS = Font(name="Arial", size=9, italic=True, color="595959")
FG = Font(name="Arial", size=10, bold=True, color="1F3864")
HF = PatternFill("solid", fgColor="1F3864")
GF = PatternFill("solid", fgColor="D9E1F2")
RED = PatternFill("solid", fgColor="FFC7CE")
GRN = PatternFill("solid", fgColor="C6EFCE")
ORG = PatternFill("solid", fgColor="FCE4D6")
WR = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
USD = '#,##0;[Red](#,##0);"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
X2 = '0.00;[Red]-0.00;"-"'


def rows(e, k):
    return [x or 0 for x in R[e]["rows"][k]]


def kpi(e, lab):
    return R[e]["kpi"].get(lab, (None, None))[1]


def annual(e, k, end=False):
    v = rows(e, k)
    return [(v[12 * y + 11] if end else sum(v[12 * y:12 * y + 12])) for y in range(10)]


# ------------------------------------------------------------------ verificaciones
CHECKS = []


def ck(nombre, aplica, fn):
    ok, det = 0, []
    for e in aplica:
        res = fn(e)
        if res is True:
            ok += 1
        else:
            det.append(f"{e}: {res}")
    CHECKS.append((nombre, len(aplica), ok, "; ".join(det) if det else ""))


todos = list(R)
tipo = lambda t: [e for e in R if R[e]["tipo"] == t]
base_rows = R[B]["rows"]


def eq(e, k, tol=1e-6):
    return max(abs(a - b) for a, b in zip(rows(e, k), rows(B, k))) < tol


ck("Sin errores de fórmula (#¡VALOR!, #¡REF!, etc.) después de recalcular", todos, lambda e: True if R[e]["err"] == 0 else f"{R[e]['err']} errores")
ck("Balance mensual cuadra (activo − pasivo − patrimonio = 0) en presupuesto y vigente", todos,
   lambda e: True if max(map(abs, rows(e, "V.CHK") + rows(e, "P.CHK"))) < 0.01 else "descuadre")


def conc(e):
    t = 0
    for i in range(N):
        t += (rows(e, "V.CFO")[i] - rows(e, "V.CAPEX")[i] + rows(e, "V.DESEMB")[i] - rows(e, "V.AMORT")[i]
              - rows(e, "V.INT")[i] * (1 if i < FIN else 0) + rows(e, "V.APORTE")[i] - rows(e, "V.DIV")[i] + rows(e, "V.AJ")[i])
    d = t - rows(e, "V.CAJA")[N - 1]
    return True if abs(d) < 0.05 else f"diferencia {d:.2f}"


ck("Caja final = suma de flujos operativos, de inversión y de financiamiento", todos, conc)
ck("El presupuesto original (bloque P) no cambia con ningún escenario (salvo X1, que lo reorganiza a propósito)", [e for e in R if R[e]["tipo"] != "particion"],
   lambda e: True if all(R[e]["rows"][k] == base_rows[k] for k in base_rows if k.startswith("P.")) else "cambió")
ck("Saldos de deuda y de línea nunca negativos", todos, lambda e: True if min(rows(e, "V.DEUDA") + rows(e, "V.LCS")) > -0.01 else "saldo negativo")
ck("Línea rotativa dentro del límite", todos, lambda e: True if max(rows(e, "V.LCS")) <= R[e]["meta"].get("lim", 0) + 0.01 else "supera el límite")


def brecha_linea(e):
    lim = R[e]["meta"].get("lim", 0)
    for i, b in enumerate(rows(e, "V.BRECHA")):
        if b > 0.5 and rows(e, "V.LCA")[i] == 1 and rows(e, "V.LCS")[i] < lim - 1:
            return f"brecha en el mes {i + 1} con línea disponible"
    return True


ck("Con línea disponible y sin agotar no hay brecha de caja", todos, brecha_linea)
ck("Dividendos pagados ≤ dividendos según política", todos, lambda e: True if all(a <= b + 0.01 for a, b in zip(rows(e, "V.DIV"), rows(e, "V.DIVD"))) else "excede")
fin = tipo("fin")
ck("Cambios de financiamiento no alteran ventas, EBITDA, flujo operativo, CAPEX ni FCFF", fin,
   lambda e: True if all(eq(e, k) for k in ["V.vtas", "V.EBITDA", "V.CFO", "V.CAPEX", "V.FCFF"]) else "cambió")
ck("Cambios de financiamiento no alteran el VAN del proyecto", fin,
   lambda e: True if abs(kpi(e, "VAN del proyecto (FCFF a WACC)") - kpi(B, "VAN del proyecto (FCFF a WACC)")) < 1e-3 else "cambió")
ck("Cuotas pagadas = monto de los préstamos (se cancelan dentro del horizonte)", [e for e in fin if "lp" in R[e]["meta"]],
   lambda e: True if abs(sum(rows(e, "V.AML")) - R[e]["meta"]["lp"]) < 0.05 else f"{sum(rows(e, 'V.AML')):.2f}")


def interes(e):
    sal, it = rows(e, "V.L1_sal"), rows(e, "V.L1_int")
    return True if all(abs(it[i] - sal[i - 1] * 0.09 / 12) < 0.01 for i in range(1, FIN) if sal[i - 1] > 0) else "interés ≠ saldo × tasa / 12"


ck("Interés mensual del préstamo 1 = saldo anterior × 9% / 12", fin, interes)


def linea_int(e):
    tasa = {"L4": 0.09, "K7": 0.0799, "K8": 0.1124}.get(e, 0.10)
    s, it = rows(e, "V.LCS"), rows(e, "V.LCI")
    return True if all(abs(it[i] - s[i - 1] * tasa / 12) < 0.01 for i in range(1, FIN)) else "interés de la línea inconsistente"


ck("Interés de la línea = saldo usado del mes anterior × tasa / 12", [e for e in fin if R[e]["meta"].get("lim")], linea_int)
ck("Línea sin renovación se cancela al vencer (K5, mes 25)", ["K5"], lambda e: True if rows(e, "V.LCS")[24] < 0.01 and rows(e, "V.LCS")[23] > 0 else "no se canceló")
dso = tipo("dso")
ck("Días de cobro no alteran el EBITDA", dso, lambda e: True if eq(e, "V.EBITDA") else "cambió")
ck("Cuentas por cobrar = ventas del mes × días / (365/12)", dso,
   lambda e: True if abs(rows(e, "V.cxc")[30] - rows(e, "V.vtas")[30] * R[e]["meta"]["dso"] / (365 / 12)) < 0.01 else "no coincide")
ck("Más días de cobro bajan el VAN del proyecto; contado lo sube", dso,
   lambda e: True if ((kpi(e, "VAN del proyecto (FCFF a WACC)") < kpi(B, "VAN del proyecto (FCFF a WACC)")) == (R[e]["meta"]["dso"] > 120)) else "dirección incorrecta")
tc = tipo("tc")
ck("El tipo de cambio no altera ventas ni materiales (en USD)", tc, lambda e: True if eq(e, "V.vtas") and eq(e, "V.MAT") else "cambió")
ck("Personal administrativo (en Gs) escala exactamente con 7.300 / TC", tc,
   lambda e: True if abs(sum(rows(e, "V.PERA")) / sum(rows(B, "V.PERA")) - 7300 / R[e]["meta"]["tc"]) < 1e-6 else "no escala")
ck("Volumen sobre la capacidad activa el control C10 (V4)", ["V4"], lambda e: True if any(c[0] == "C10" and c[3] == "REVISAR" for c in R[e]["ctl"]) else "no alertó")
ck("Caja bajo el mínimo activa el control C13", [e for e in R if max(rows(e, "V.BRECHA")) > 0.5],
   lambda e: True if any(c[0] == "C13" and c[3] == "REVISAR" for c in R[e]["ctl"]) else "no alertó")
ck("Inicio atrasado: sin ventas antes del nuevo mes de inicio (P1: mes 7)", ["P1"], lambda e: True if sum(rows(e, "V.vtas")[:6]) == 0 and rows(e, "V.vtas")[6] > 0 else "ventas antes")
ck("Selector de escenario aplica el atraso de 3 meses (P2)", ["P2"], lambda e: True if sum(rows(e, "V.vtas")[:3]) == 0 and rows(e, "V.vtas")[3] > 0 else "no aplicó")


def kpis_iguales(e, ref, tol=0.01):
    malos = []
    for lab, (p, v) in ref.items():
        x = R[e]["kpi"].get(lab, (None, None))[1]
        if isinstance(v, (int, float)) and isinstance(x, (int, float)):
            if abs(v - x) > (tol if abs(v) > 10 else 1e-6):
                malos.append(f"{lab}: {v} vs {x}")
        elif v != x:
            malos.append(f"{lab}: {v} vs {x}")
    return True if not malos else "; ".join(malos[:3])


part = tipo("particion")
ck("Partición en los 25 espacios: todos los indicadores iguales a la base", part, lambda e: kpis_iguales(e, R[B]["kpi"]))
ck("Partición: ventas, materiales, EBITDA y caja mes a mes iguales a la base", part,
   lambda e: True if all(max(abs(a - b) for a, b in zip(rows(e, k), rows(B, k))) < 0.01 for k in ["V.vtas", "V.MAT", "V.EBITDA", "V.CAJA", "P.vtas", "P.MAT"]) else "difiere")
ck("Partición: sin alertas de capacidad ni de precio (C10, C11 en OK)", part,
   lambda e: True if all(c[3] == "OK" for c in R[e]["ctl"] if c[0] in ("C10", "C11")) else "alerta")
nuevo = tipo("nuevo")
ck("Productos nuevos 3, 13 y 25 (solo vigente): Δ ventas = Σ volumen × precio × índice; presupuesto sin cambio", nuevo,
   lambda e: True if abs(sum(rows(e, "V.vtas")) - sum(rows(B, "V.vtas")) - R[e]["meta"]["dv"]) < 0.01 and eq(e, "P.vtas") else "no coincide")
ck("Productos nuevos 3, 13 y 25: Δ materiales = Σ volumen × (MP × índice + insumos); presupuesto sin cambio", nuevo,
   lambda e: True if abs(sum(rows(e, "V.MAT")) - sum(rows(B, "V.MAT")) - R[e]["meta"]["dm"]) < 0.01 and eq(e, "P.MAT") else "no coincide")
sinp = tipo("sinprecio")
ck("Producto con volumen y sin precio activa C11 y no suma ventas", sinp,
   lambda e: True if any(c[0] == "C11" and c[3] == "REVISAR" for c in R[e]["ctl"]) and eq(e, "V.vtas") else "no alertó")
if REF:
    comunes = [e for e in REF if e in R]
    ck("Regresión: cada escenario da los mismos indicadores que la corrida de referencia (12 productos)", comunes, lambda e: kpis_iguales(e, REF[e]["kpi"]))
    ck("Regresión: caja, deuda y FCFF mes a mes iguales a la referencia", comunes,
       lambda e: True if all(max(abs((a or 0) - (b or 0)) for a, b in zip(R[e]["rows"][k], REF[e]["rows"][k])) < 0.01 for k in ["V.CAJA", "V.DEUDA", "V.FCFF", "V.EBITDA"]) else "difiere")

# ------------------------------------------------------------------ libro
wb = Workbook()
ws = wb.active
ws.title = "Resumen"
ws["A1"] = f"{MODELO} — batería de escenarios (resultados del modelo, vigente)"
ws["A1"].font = FT
ws["A2"] = (f"Valores fijos de las corridas del {FECHA} sobre {LIBRO} (supuestos Oct-2025, sin reales). "
            "Cada escenario modifica solo entradas vigentes; el presupuesto original queda como referencia. USD.")
ws["A2"].font = FS
cols = [("Id", 6, None), ("Escenario", 34, None), ("Qué cambia", 46, None),
        ("Ventas 10 años", 13, USD), ("EBITDA 10 años", 13, USD), ("Resultado neto 10 años", 13, USD),
        ("VAN proyecto (9,49%)", 13, USD), ("Δ VAN proyecto vs base", 13, USD), ("TIR proyecto", 9, PCT),
        ("VAN socios (16,32%)", 13, USD), ("Δ VAN socios vs base", 13, USD), ("TIR socios", 9, PCT),
        ("Necesidad de fondos antes de financiar", 14, USD), ("Aporte adicional necesario (caja negativa)", 14, USD), ("Meses bajo caja mínima", 9, '0'),
        ("Deuda máxima", 13, USD), ("Línea rotativa: uso máximo", 13, USD), ("Intereses totales", 12, USD),
        ("Cancelación de la deuda (años)", 11, X2), ("DSCR mínimo desde año 2", 10, X2), ("Recupero (años)", 10, X2)]
L = {"vt": "Ventas totales del horizonte", "et": "EBITDA total del horizonte", "nt": "Resultado neto total", "van": "VAN del proyecto (FCFF a WACC)",
     "tir": "TIR del proyecto (anual)", "vane": "VAN de los socios (a Ke)", "tire": "TIR de los socios (anual)",
     "nec": "Necesidad máxima de fondos antes de financiamiento", "ap": "Aporte adicional necesario para no tener caja negativa", "br": "Meses con brecha",
     "deu": "Deuda financiera máxima", "lc": "Línea rotativa: uso máximo", "int": "Intereses, IVA y comisiones totales",
     "can": "Cancelación total de la deuda (años desde el mes 1)", "dscr": "DSCR mínimo desde el 2.º año de operación (flujo operativo / servicio de deuda)",
     "pb": "Recupero de la inversión del proyecto (años)"}
r = 4
for i, (h, w, _) in enumerate(cols, 1):
    c = ws.cell(r, i, h)
    c.font, c.fill, c.alignment = FH, HF, CEN
    ws.column_dimensions[CL(i)].width = w
ws.row_dimensions[r].height = 48
r += 1
first = r
grupo = None
for e, v in R.items():
    if v["grupo"] != grupo:
        grupo = v["grupo"]
        ws.cell(r, 1, grupo).font = FG
        for c in range(1, len(cols) + 1):
            ws.cell(r, c).fill = GF
        r += 1
    vals = [e, v["nombre"], v["cambios"], kpi(e, L["vt"]), kpi(e, L["et"]), kpi(e, L["nt"]), kpi(e, L["van"]),
            kpi(e, L["van"]) - kpi(B, L["van"]), kpi(e, L["tir"]), kpi(e, L["vane"]), kpi(e, L["vane"]) - kpi(B, L["vane"]), kpi(e, L["tire"]),
            kpi(e, L["nec"]), kpi(e, L["ap"]), kpi(e, L["br"]), kpi(e, L["deu"]), kpi(e, L["lc"]), kpi(e, L["int"]), kpi(e, L["can"]),
            kpi(e, L["dscr"]), kpi(e, L["pb"])]
    for i, (x, (_, _, fmt)) in enumerate(zip(vals, cols), 1):
        c = ws.cell(r, i, x)
        c.font = FB if e == B else F
        if fmt and not isinstance(x, str):
            c.number_format = fmt
        if i == 3:
            c.alignment = WR
    r += 1
last = r - 1
ws.conditional_formatting.add(f"N{first}:N{last}", CellIsRule(operator="greaterThan", formula=["0.5"], fill=ORG))
ws.conditional_formatting.add(f"T{first}:T{last}", CellIsRule(operator="lessThan", formula=["1"], fill=ORG))
ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(operator="lessThan", formula=["0"], fill=RED))
ws.conditional_formatting.add(f"J{first}:J{last}", CellIsRule(operator="lessThan", formula=["0"], fill=RED))
r += 1
notas = [
    "Lectura: VAN/TIR del proyecto = flujo libre (FCFF) sin financiamiento, a WACC 9,49% (histórico del plan). VAN/TIR de los socios = aportes, dividendos y valor final, a Ke 16,32%.",
    "«Aporte adicional necesario» = máximo saldo de caja negativa no cubierto por los préstamos/aportes cargados. En el VAN de los socios se considera aporte implícito (costo Ke) que se devuelve al recuperarse la caja.",
    "Naranja: necesidad de fondos no cubierta o DSCR < 1 (la operación no alcanza a pagar las cuotas ese año). Rojo: VAN negativo.",
    "Tasas: LP 9% nominal (plan). CP 10% ilustrativa; sensibilidad con 7,99% (tasa activa promedio ME, BCP mar-2026) y 11,24% (tope ME, BCP jul-2026). Las tasas reales dependen de la oferta bancaria.",
    "La caja positiva no genera intereses en el modelo: el exceso de financiamiento solo suma costo.",
]
for t in notas:
    ws.cell(r, 2, t).font = FS
    r += 1
ws.freeze_panes = "D5"

# ------------------------------------------------------------------ anuales
for nombre, k, end in [("Caja anual", "V.CAJA", True), ("Deuda anual", "V.DEUDA", True), ("Linea anual", "V.LCS", True),
                       ("EBITDA anual", "V.EBITDA", False), ("FCFF anual", "V.FCFF", False)]:
    w = wb.create_sheet(nombre)
    w["A1"] = f"{nombre.replace('Linea', 'Línea rotativa')} por escenario (USD, {'saldo al cierre del año' if end else 'suma del año'})"
    w["A1"].font = FT
    hdr = ["Id", "Escenario"] + [f"Año {y}" for y in range(1, 11)]
    for i, h in enumerate(hdr, 1):
        c = w.cell(3, i, h)
        c.font, c.fill, c.alignment = FH, HF, CEN
        w.column_dimensions[CL(i)].width = 6 if i == 1 else (40 if i == 2 else 12)
    rr = 4
    for e, v in R.items():
        w.cell(rr, 1, e).font = F
        w.cell(rr, 2, v["nombre"]).font = F
        for y, x in enumerate(annual(e, k, end), 3):
            c = w.cell(rr, y, x)
            c.number_format, c.font = USD, (FB if e == B else F)
        rr += 1
    if k == "V.CAJA":
        w.conditional_formatting.add(f"C4:L{rr - 1}", CellIsRule(operator="lessThan", formula=["0"], fill=RED))
    w.freeze_panes = "C4"

# ------------------------------------------------------------------ verificaciones
w = wb.create_sheet("Verificaciones")
w["A1"] = f"Verificaciones automáticas del motor sobre los {len(R)} escenarios"
w["A1"].font = FT
for i, (h, wd) in enumerate([("#", 5), ("Verificación", 80), ("Escenarios", 11), ("OK", 8), ("Estado", 10), ("Detalle de fallas", 50)], 1):
    c = w.cell(3, i, h)
    c.font, c.fill, c.alignment = FH, HF, CEN
    w.column_dimensions[CL(i)].width = wd
rr = 4
for i, (n, a, ok, det) in enumerate(CHECKS, 1):
    for j, x in enumerate([i, n, a, ok, "OK" if ok == a else "FALLA", det], 1):
        c = w.cell(rr, j, x)
        c.font = F
    w.cell(rr, 5).fill = GRN if ok == a else RED
    rr += 1
tot = sum(c[1] for c in CHECKS)
okt = sum(c[2] for c in CHECKS)
w.cell(rr + 1, 2, f"Total: {okt} de {tot} comprobaciones correctas ({len(CHECKS)} tipos de verificación).").font = FB

# ------------------------------------------------------------------ supuestos
w = wb.create_sheet("Supuestos")
w["A1"] = "Supuestos de las pruebas"
w["A1"].font = FT
w.column_dimensions["A"].width = 34
w.column_dimensions["B"].width = 110
sup = [
    ("Base", "Plan Ecostar Oct-2025 cargado en la plantilla v2: ventas USD 81,2 M en 10 años, CAPEX USD 755 mil, préstamo LP USD 2,79 M al 9% (6 meses de gracia + 54 cuotas), aportes USD 1,14 M."),
    ("Escenarios", "Se modifican solo entradas vigentes, como haría el usuario; el presupuesto original no cambia. Sin datos reales (todo es proyección)."),
    ("Más / menos financiamiento", "± USD 1,0 M sobre el préstamo LP (3,8 M / 1,8 M), mismas condiciones. Magnitud ilustrativa."),
    ("Cobranza", "180 días (vs 120 del plan) o contado (0 días). Magnitudes ilustrativas."),
    ("Costos", "Costos variables por unidad +10/20/30%; caso extremo: además salarios +30% y costos fijos ×2. Ilustrativo."),
    ("Ventas", "Volumen ±10/20/30%; «dentro de la capacidad» lleva los años 1-2 a la capacidad inicial (15.000 gabinetes / 90.000 cajas)."),
    ("Plan desfavorable", "Volumen −30%, precio −5%, inicio 6 meses después, CAPEX +15%. Ilustrativo."),
    ("Tasas de financiamiento", "LP 9% (plan). CP 10% ilustrativa. Referencias BCP: tasa activa promedio ponderada ME 7,99% (mar-2026); tope legal ME 11,24% (jul-2026). "
                                "Fuente: BCP, Indicadores Financieros e informes de tasas máximas (consulta 01/10/2026; no se pudo abrir el sitio del BCP desde este entorno)."),
    ("Tipo de cambio", "Base Gs 7.300 (plan). Sube a 8.760 (+20%); baja a 6.145 (nivel del archivo de Alianza del Acero, jun-2026) y 5.840 (−20%). Ventas y materiales en USD; nómina y parte de los fijos en Gs."),
    ("Línea rotativa", "Se usa automáticamente cuando la caja cae bajo el mínimo (3 días de ventas) y se devuelve con el excedente; los dividendos esperan a que se cancele. "
                       "Al vencer se paga todo el saldo. Supone renovación garantizada hasta el vencimiento cargado."),
    ("Estructura de 25 productos", "X1-X3 son pruebas técnicas de la ampliación a 25 productos con datos de prueba (no comerciales): partición del volumen actual "
                                   "en los 25 espacios (debe dar lo mismo que la base), productos nuevos 3, 13 y 25 (sumas exactas) y producto sin precio (alerta C11)."),
    ("Limitaciones", "La caja positiva no genera intereses. No se modelan comisiones de apertura ni garantías. El volumen sobre capacidad no se limita (solo se alerta). "
                     "Los resultados usan supuestos de 2025 y no acreditan la situación actual."),
]
for i, (a, b) in enumerate(sup, 3):
    w.cell(i, 1, a).font = FB
    c = w.cell(i, 2, b)
    c.font, c.alignment = F, WR
    w.row_dimensions[i].height = 30
for s in wb.worksheets:
    s.sheet_view.zoomScale = 90
wb.save(OUT)
print(f"Guardado {OUT}: {okt}/{tot} comprobaciones OK")
for c in CHECKS:
    if c[1] != c[2]:
        print("FALLA:", c)
