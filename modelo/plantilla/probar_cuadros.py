# -*- coding: utf-8 -*-
"""
Pruebas de las hojas de cuadros (estructura del plan de negocio).
Uso: python3 probar_cuadros.py libro.xlsx recalc.py carpeta

Recalcula 7 casos en LibreOffice y verifica que cada cuadro cuadre con el motor y con Resultados:
fuentes = usos (cuadro 18), totales iguales a Resultados, costos = ventas − utilidad, VAN y TIR, saldos al cierre,
filas de ajuste en 0 sin reales ni detalle mensual, selector de versión, años de operación con escenario e implantación,
y el método del plan aplicado a sus propios flujos (VAN y TIR publicados). Guarda los resultados en cuadros.pkl.
"""
import sys, os, json, subprocess, pickle
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook

BASE, RECALC, OUT = sys.argv[1:4]
os.makedirs(OUT, exist_ok=True)
FC = 6
REALES = {"TC": [7400] * 6, "U1": [600] * 6, "U2": [4000] * 6, "S1": [162000] * 6, "S2": [92000] * 6, "MAT": [120000] * 6, "OVAR": [30000] * 6,
          "PERP": [300000000] * 6, "PERA": [60000000] * 6, "FIJP": [15000] * 6, "ADM": [25000] * 6, "TRIB": [2540] * 6,
          "CAPEX": [600000, 0, 0, 30000, 0, 0], "DESEMB": [2791587, 0, 0, 0, 0, 0], "INT": [23026] * 6, "AMORT": [0] * 6,
          "APORTE": [767782, 0, 0, 0, 0, 0], "DIV": [0] * 6}
SALDOS = {"S_CAJA": 1500000, "S_CXC": 900000, "S_INV": 500000, "S_IVA": 0, "S_PROV": 100000, "S_DEUDA": 2791587}


def rows_by_code(ws):
    return {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value is not None}


def name_cell(wb, name):
    sh, rg = wb.defined_names[name].attr_text.split("!")
    return wb[sh.strip("'")], rg.replace("$", "")


def setn(wb, name, val):
    ws, rg = name_cell(wb, name)
    ws[rg] = val


def reales(wb):
    ws = wb["Reales"]; rr = rows_by_code(ws)
    for k, v in REALES.items():
        for m, x in enumerate(v, 1):
            ws.cell(rr[k], FC + m - 1).value = x
    for k, v in SALDOS.items():
        ws.cell(rr[k], FC + 5).value = v
    inv = wb["Inversion"]
    for r in range(6, 17):
        inv.cell(r, 7).value = inv.cell(r, 4).value
        inv.cell(r, 8).value = inv.cell(r, 4).value
    inv.cell(17, 7).value = 35833.91
    inv.cell(17, 8).value = 35833.91
    setn(wb, "n_Corte", 6)


def mp_detalle(wb):
    ws = wb["Detalle_24m"]; r = rows_by_code(ws)["MP1"]
    for m in range(1, 25):
        ws.cell(r, 5 + m).value = 90


CASES = {
    "base": [],
    "pres": [lambda w: w["Cuadros"].__setitem__("C5", "Presupuesto original")],
    "cons": [lambda w: setn(w, "n_Escenario", "Conservador")],
    "reales_cons": [reales, lambda w: setn(w, "n_Escenario", "Conservador")],
    "mp": [mp_detalle],
    "impl": [lambda w: setn(w, "o_Inicio", 7)],
    "impl_pres": [lambda w: setn(w, "o_Inicio", 7), lambda w: w["Cuadros"].__setitem__("C5", "Presupuesto original")],
}
HOJAS = ["Cuadros", "C01 Inversión", "C07 Producción", "C08 Ingresos", "C09 Materiales", "C10a Costos", "C10b Costo unitario", "C11 Rubros",
         "C13 Deuda e IVA", "C15 Resultados", "C16 Capital trabajo", "C17 Equilibrio", "C18 Fondos", "C19 VAN-TIR", "C20 Depreciación", "RRHH"]


def leer_cuadro(ws):
    """{bloque: {clave: [año 0..10, total]}} y filas sueltas por etiqueta."""
    blk, out, lab = None, {"A": {}, "B": {}, "C": {}}, {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a[:3] in ("A. ", "B. ", "C. "):
            blk = a[0]
            continue
        b = ws.cell(r, 2).value
        if isinstance(a, str) and blk and "." in a:
            out[blk][a] = [ws.cell(r, c).value for c in range(4, 16)]
        if isinstance(b, str):
            lab[b] = [ws.cell(r, c).value for c in range(3, 9)]
    return out, lab


def run(item):
    name, edits = item
    wb = load_workbook(BASE)
    for f in edits:
        f(wb)
    p = os.path.join(OUT, f"q_{name}.xlsx"); wb.save(p)
    out = subprocess.run(["python3", RECALC, p, "900"], capture_output=True, text=True).stdout
    j = json.loads(out[out.index("{"):])
    w = load_workbook(p, data_only=True)
    c = w["Calculo"]
    rows = {c.cell(r, 1).value: [c.cell(r, k).value for k in range(5, 150)] for r in range(1, c.max_row + 1) if isinstance(c.cell(r, 1).value, str)}
    res = w["Resultados"]
    kpi = {res.cell(r, 2).value: (res.cell(r, 3).value, res.cell(r, 4).value) for r in range(6, 31)}
    cu = {h: leer_cuadro(w[h]) for h in HOJAS}
    return name, dict(err=j.get("total_errors"), errs=j.get("error_summary"), rows=rows, kpi=kpi, cu=cu)


with ThreadPoolExecutor(int(os.environ.get("HILOS", 4))) as ex:
    R = dict(ex.map(run, CASES.items()))

RES = []


def chk(lab, ok):
    RES.append((lab, bool(ok)))
    print(("OK   " if ok else "FALLA"), lab)


num = lambda v: v if isinstance(v, (int, float)) else 0
near = lambda a, b, t=0.01: abs(num(a) - num(b)) <= t


def A(n, h, k):
    return R[n]["cu"][h][0]["A"][k]


def anual(n, X, key, y, kind="s"):
    """Suma (s) o saldo al cierre (e) del año de operación y en el motor (año 0 = implantación)."""
    rr = R[n]["rows"]
    vals, aop, pre, m = rr[f"{X}.{key}"][1:], rr[f"{X}.aop"][1:], rr[f"{X}.pre"][1:], list(range(1, 145))
    sel = [i for i in range(144) if (pre[i] == 1 if y == 0 else aop[i] == y)]
    if not sel:
        return 0
    return sum(num(vals[i]) for i in sel) if kind == "s" else num(vals[max(sel)])


for n in CASES:
    X = "P" if n in ("pres", "impl_pres") else "V"
    col = 0 if X == "P" else 1
    kp = R[n]["kpi"]
    chk(f"{n}: sin errores de fórmula", R[n]["err"] == 0)
    chk(f"{n}: cuadro 18 — control fuentes − usos − caja = 0 en todos los años",
        all(near(v, 0) for v in A(n, "C18 Fondos", "c18.chk")[:11]))
    chk(f"{n}: cuadro 18 — caja al cierre de cada año = motor",
        all(near(A(n, "C18 Fondos", "c18.sal")[y], anual(n, X, "CAJA", y, "e"), 0.01) for y in range(0, 11)))
    for k, lab in [("c15.vtas", "Ventas totales del horizonte"), ("c15.ebitda", "EBITDA total del horizonte"), ("c15.ni", "Resultado neto total")]:
        chk(f"{n}: cuadro 15 — {lab} = Resultados", near(A(n, "C15 Resultados", k)[11], kp[lab][col], 0.01))
    chk(f"{n}: cuadro 15 — año por año = motor (EBITDA)", all(near(A(n, "C15 Resultados", "c15.ebitda")[y], anual(n, X, "EBITDA", y), 0.01) for y in range(11)))
    v, t, ni = A(n, "C10a Costos", "c10.tot"), A(n, "C15 Resultados", "c15.vtas"), A(n, "C15 Resultados", "c15.ni")
    ire = A(n, "C15 Resultados", "c15.ire")
    chk(f"{n}: cuadro 10a — ventas − costos totales = utilidad + IRE (cada año)", all(near(num(t[y]) - num(v[y]), num(ni[y]) + num(ire[y]), 0.02) for y in range(11)))
    chk(f"{n}: cuadro 11 — rubros + ajuste = costos totales (cada año)",
        all(near(sum(num(A(n, "C11 Rubros", f"c11.{k}")[y]) for k in ["mo", "mp", "ot", "imp", "fin", "dep", "aj"]), A(n, "C11 Rubros", "c11.tot")[y], 0.02) for y in range(11)))
    chk(f"{n}: cuadro 11 = cuadro 10a (costos totales)", all(near(A(n, "C11 Rubros", "c11.tot")[y], v[y], 0.02) for y in range(11)))
    chk(f"{n}: cuadro 19 — FCFF descontado total = VAN del proyecto", near(A(n, "C19 VAN-TIR", "c19.fcffd")[11], kp["VAN del proyecto (FCFF a WACC)"][col], 0.05))
    chk(f"{n}: cuadro 19 — flujo de los socios descontado = VAN de los socios", near(A(n, "C19 VAN-TIR", "c19.eqd")[11], kp["VAN de los socios (a Ke)"][col], 0.05))
    chk(f"{n}: cuadro 16 — capital de trabajo neto al cierre = motor (años 1-10)",
        all(near(A(n, "C16 Capital trabajo", "c16.ctn")[y], anual(n, X, "ctn", y, "e"), 0.01) for y in range(1, 11)))
    chk(f"{n}: cuadro 16 — bruto − proveedores = neto (cada año)",
        all(near(num(A(n, "C16 Capital trabajo", "c16.bruto")[y]) - num(A(n, "C16 Capital trabajo", "c16.pas")[y]), A(n, "C16 Capital trabajo", "c16.ctn")[y], 0.02)
            for y in range(1, 11)))
    chk(f"{n}: cuadro 20 — depreciación por ítem + ajuste = motor", all(near(num(A(n, "C20 Depreciación", "c20.tot")[y]) + num(A(n, "C20 Depreciación", "c20.aj")[y]),
                                                                        anual(n, X, "DEP", y), 0.01) for y in range(11)))
    chk(f"{n}: cuadro 13 — saldo de préstamos al cierre del año 10 = 0", near(A(n, "C13 Deuda e IVA", "c13.sal")[10], 0, 0.01))
    chk(f"{n}: cuadro 8 — ingresos por producto suman el total", all(near(sum(num(A(n, "C08 Ingresos", f"c08.s.{i}")[y]) for i in range(1, 26)),
                                                                        A(n, "C08 Ingresos", "c08.vtas")[y], 0.02) for y in range(11)))
    chk(f"{n}: cuadro 10b — costo total asignado por producto suma el total",
        all(near(sum(num(A(n, "C10b Costo unitario", f"c10b.ct.{i}")[y]) for i in range(1, 26)), A(n, "C10b Costo unitario", "c10b.tot")[y], 0.05)
            for y in range(1, 11)))
    if n in ("base", "pres", "cons", "impl", "impl_pres"):
        aj = [("C09 Materiales", "c09.mpaj"), ("C09 Materiales", "c09.insaj"), ("C09 Materiales", "c09.ovaj"), ("C10a Costos", "c10.ajp"),
              ("C10a Costos", "c10.aja"), ("C10a Costos", "c10.fcaj"), ("C11 Rubros", "c11.aj"), ("C17 Equilibrio", "c17.faj"), ("C20 Depreciación", "c20.aj")]
        chk(f"{n}: filas de ajuste en 0 (sin reales ni detalle mensual)", all(near(A(n, h, k)[y], 0, 0.05) for h, k in aj for y in range(12)))

B = R["base"]["cu"]
chk("base: método del plan con los flujos publicados = VAN publicado (4.588.736)", near(B["C19 VAN-TIR"][0]["B"]["c19.p_act"][11], 4588736, 2))
lab = B["C19 VAN-TIR"][1]
tir_b = lab["TIR — método del plan; plan: recalculada con los flujos publicados"][2]
chk("base: método del plan con los flujos publicados = TIR publicada (22,61%)", near(tir_b, 0.2261, 0.0001))
chk("base: cuadro 7 — utilización de gabinetes del año 1 = 80%", near(A("base", "C07 Producción", "c07.ut.1")[1], 0.8, 1e-9))
chk("base: cuadro 8 — precio de gabinetes del año 3 = 270 × 1,01² (275,43)", near(A("base", "C08 Ingresos", "c08.p.1")[3], 270 * 1.01 ** 2, 1e-6))
chk("base: RRHH — 65 personas en el año 1 (43 directos + 22 indirectos y administración)",
    near(A("base", "RRHH", "rh.ntot")[1], 65) and near(A("base", "RRHH", "rh.ndir")[1], 43) and near(A("base", "RRHH", "rh.nia")[1], 22))
chk("base: cuadro 12 — estructura de costos suma 100%", all(near(sum(num(A("base", "C11 Rubros", f"c12.{k}")[y]) for k in ["mo", "mp", "ot", "imp", "fin", "dep"]), 1, 1e-6)
                                                          for y in range(1, 11)))
chk("base: cuadro 1 — inversión fija inicial (pagada hasta el año 1) = 655.044",
    near(R["base"]["cu"]["C01 Inversión"][1]["   pagada hasta el año 1 de operación (inversión inicial)"][1], 655044.09, 0.01))
idx = B["Cuadros"][1]
chk("base: índice — diferencia de EBITDA de 10 años frente al plan = −1.044.506", near(idx["EBITDA de los 10 años"][2], -1044506, 1))
chk("pres: versión «Presupuesto original» = base (vigente sin cambios)", all(near(A("pres", "C15 Resultados", "c15.ebitda")[y], A("base", "C15 Resultados", "c15.ebitda")[y]) for y in range(12)))
chk("cons: año 0 sin ventas y con costos de implantación (atraso de 3 meses)",
    near(A("cons", "C08 Ingresos", "c08.vtas")[0], 0) and num(A("cons", "C10a Costos", "c10.tot")[0]) > 0)
chk("cons: el año 1 de operación empieza en el mes 4 (ventas del año 1 = motor)", near(A("cons", "C08 Ingresos", "c08.vtas")[1], anual("cons", "V", "vtas", 1), 0.01))
chk("cons: el presupuesto original no cambia (Resultados, columna presupuesto)",
    all(near(R["cons"]["kpi"][k][0], R["base"]["kpi"][k][0]) for k in ["Ventas totales del horizonte", "EBITDA total del horizonte", "VAN del proyecto (FCFF a WACC)"]))
est = R["reales_cons"]["cu"]["C15 Resultados"][1]["Estado del año"]
chk("reales_cons: estado del año 1 = «Real+Proy.» y del año 2 = «Proyección»", est[2] == "Real+Proy." and est[3] == "Proyección")
mpaj = A("mp", "C09 Materiales", "c09.mpaj")
chk("mp: el detalle de materia prima aparece en el ajuste del cuadro 9 (año 1: 12.000 × 7,35 = 88.200; año 3: 0)",
    near(mpaj[1], 12000 * 7.35, 0.01) and near(mpaj[3], 0, 0.01))
chk("impl: año 0 = 6 meses de implantación (sin ventas, con costos e inversión)",
    near(A("impl", "C08 Ingresos", "c08.vtas")[0], 0) and num(A("impl", "C10a Costos", "c10.tot")[0]) > 0)
chk("impl: los años de operación 1-10 suman las ventas del horizonte", near(sum(num(x) for x in A("impl", "C08 Ingresos", "c08.vtas")[:11]),
                                                                             R["impl"]["kpi"]["Ventas totales del horizonte"][1], 0.01))
pickle.dump({"R": {k: {kk: vv for kk, vv in v.items() if kk != "rows"} for k, v in R.items()}, "RES": RES}, open(os.path.join(OUT, "cuadros.pkl"), "wb"))
print(f"\n{sum(ok for _, ok in RES)} de {len(RES)} comprobaciones OK")
