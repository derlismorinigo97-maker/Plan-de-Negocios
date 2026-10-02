# -*- coding: utf-8 -*-
"""
Plantilla común de modelo financiero para maquilas (v2 simplificada).

Un archivo independiente por negocio, misma estructura:
  Inicio · Supuestos · Productos · Costos · Inversion · Reales · Calculo · Resultados · Control

- Entradas solo en Inicio, Supuestos, Productos, Costos, Inversion y Reales (celdas amarillas, texto azul).
- Cada dato se carga una vez: la columna «Vigente» hereda el «Presupuesto original» (gris) hasta que se sobrescribe.
- Calculo es un motor mensual único que se ejecuta dos veces con las mismas fórmulas:
  P = presupuesto original (sin reales) y V = vigente (reales hasta el mes de corte + proyección con escenario).
- Resultados y Control se calculan automáticamente.

Uso:
  python3 generar_plantilla.py ecostar  salida.xlsx    # carga de Ecostar
  python3 generar_plantilla.py vacia    salida.xlsx    # plantilla en blanco
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

MODO = sys.argv[1] if len(sys.argv) > 1 else "ecostar"
OUT = sys.argv[2] if len(sys.argv) > 2 else "Modelo_Maquila_v2.xlsx"
if MODO == "ecostar":
    from datos_ecostar import DATOS as D
else:
    D = {"negocio": None, "socios": None, "etapa": None, "producto_desc": None, "fecha_mes1": None, "corte": None,
         "inicio_real": None, "escenario": "Vigente", "params": {"Anios": (10, "")}, "anual": {}, "productos": [],
         "volumen": [], "personal": [], "fijos": [], "capex": [], "capex_nota": "", "prestamos": [], "aportes": [],
         "publicado": {}, "pendientes": []}

# ---------------------------------------------------------------------------- dimensiones
N = 144                     # meses de la rejilla (12 años: implantación + 10 años de operación)
FCOL = 6                    # columna del mes 1 en Reales y Calculo (F)
NP, NPER, NFIJ, NCAP, NAP, NL = 25, 30, 25, 20, 10, 3   # productos individuales, puestos, costos fijos, CAPEX, aportes, préstamos
NY = 10                     # años de supuestos
NYR = 12                    # años del resumen


def mc(m):
    return CL(FCOL + m - 1)


M1, MN = mc(1), mc(N)

# ---------------------------------------------------------------------------- estilos
FONT = "Arial"
F = Font(name=FONT, size=9)
F_IN = Font(name=FONT, size=9, color="0000FF")
F_INH = Font(name=FONT, size=9, color="808080", italic=True)
F_LINK = Font(name=FONT, size=9, color="008000")
F_B = Font(name=FONT, size=9, bold=True)
F_H = Font(name=FONT, size=9, bold=True, color="FFFFFF")
F_T = Font(name=FONT, size=13, bold=True, color="1F3864")
F_S = Font(name=FONT, size=9, italic=True, color="595959")
F_SEC = Font(name=FONT, size=10, bold=True, color="1F3864")
F_AL = Font(name=FONT, size=9, bold=True, color="C00000")
FL_IN = PatternFill("solid", fgColor="FFF2CC")
FL_INH = PatternFill("solid", fgColor="F2F2F2")
FL_H = PatternFill("solid", fgColor="1F3864")
FL_SEC = PatternFill("solid", fgColor="D9E1F2")
FL_TOT = PatternFill("solid", fgColor="EDEDED")
FL_P = PatternFill("solid", fgColor="FCE4D6")     # presupuesto original
FL_V = PatternFill("solid", fgColor="E2EFDA")     # vigente
FL_OK = PatternFill("solid", fgColor="C6EFCE")
FL_BAD = PatternFill("solid", fgColor="FFC7CE")
NF_USD = '#,##0;(#,##0);"-"'
NF_USD2 = '#,##0.00;(#,##0.00);"-"'
NF_USD4 = '#,##0.0000;(#,##0.0000);"-"'
NF_PCT = '0.0%;(0.0%);"-"'
NF_PCT2 = '0.00%;(0.00%);"-"'
NF_INT = '#,##0;(#,##0);"-"'
NF_N = '0'
NF_X = '0.00"x";(0.00"x");"-"'
NF_DATE = 'dd/mm/yyyy'
NF_MES = 'mmm-yy'
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
wb.remove(wb.active)
SHEETS = ["Inicio", "Supuestos", "Productos", "Costos", "Inversion", "Reales", "Calculo", "Resultados", "Control"]
WS = {s: wb.create_sheet(s) for s in SHEETS}
CUR = [None]


# Criterio: no usar decimales dentro de criterios de texto (COUNTIF(..., ">0.5")): Excel los interpreta según la
# configuración regional (coma decimal en es-PY) y devuelve 0. Usar comparaciones numéricas (SUMPRODUCT(--(rango>0.5))).


def nm(name, sheet, ref):
    """Nombre de libro: ref puede ser celda (D5) o rango (D5:M5)."""
    parts = ref.split(":")
    absref = ":".join("$" + "".join(c for c in p if c.isalpha()) + "$" + "".join(c for c in p if c.isdigit()) for p in parts)
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!{absref}")


def put(ws, r, c, v, font=None, fill=None, fmt=None, al=None, note=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font or F
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if al:
        cell.alignment = al
    if note:
        cell.comment = Comment(note, "Modelo")
    return cell


def inp(ws, r, c, v=None, fmt=None):
    return put(ws, r, c, v, font=F_IN, fill=FL_IN, fmt=fmt)


def inh(ws, r, c, formula, fmt=None):
    """Celda vigente que hereda del presupuesto: gris hasta que se sobrescribe."""
    return put(ws, r, c, formula, font=F_INH, fill=FL_INH, fmt=fmt)


def title(ws, t, sub=None):
    put(ws, 1, 1, t, font=F_T)
    if sub:
        put(ws, 2, 1, sub, font=F_S)


def section(ws, r, t, ncol=16):
    for c in range(1, ncol + 1):
        ws.cell(row=r, column=c).fill = FL_SEC
    put(ws, r, 1, t, font=F_SEC, fill=FL_SEC)


def header(ws, r, labels, c0=1, h=None):
    for i, l in enumerate(labels):
        put(ws, r, c0 + i, l, font=F_H, fill=FL_H, al=CEN)
    if h:
        ws.row_dimensions[r].height = h


def dv(ws, rng, opts):
    d = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
    ws.add_data_validation(d)
    d.add(rng)


def widths(ws, w):
    for k, v in w.items():
        ws.column_dimensions[k].width = v


def P(key, i=None):
    """Valor de un parámetro del diccionario de datos."""
    v = D.get("params", {}).get(key)
    return v[0] if v else None


# ============================================================================ INICIO
def build_inicio():
    ws = WS["Inicio"]
    widths(ws, {"A": 3, "B": 46, "C": 34, "D": 90})
    put(ws, 1, 1, '="Modelo financiero de maquila — "&IF(n_Negocio="","(completar nombre)",n_Negocio)', font=F_T)
    put(ws, 2, 1, "Plantilla común v2: un archivo por negocio, misma estructura. Presupuesto original, ejecución real y proyección vigente separados.", font=F_S)
    r = 4
    section(ws, r, "1. Configuración del negocio (completar)", 4); r += 1
    header(ws, r, ["", "Dato", "Valor", "Instrucción"], 1); r += 1
    rows = [
        ("n_Negocio", "Nombre del negocio", D.get("negocio"), None, "Texto libre."),
        ("n_Socios", "Socios / empresas", D.get("socios"), None, "Texto libre."),
        ("n_Producto", "Actividad / productos", D.get("producto_desc"), None, "Texto libre."),
        ("n_Etapa", "Etapa del negocio", D.get("etapa"), None, "Implantación (aún sin operar) u Operación."),
        ("n_Fecha1", "Fecha del mes 1 del modelo", D.get("fecha_mes1"), NF_DATE,
         "Primer mes de implantación u operación. Se recomienda que sea enero para que los años coincidan con años calendario. Vacío = meses relativos."),
        ("n_Corte", "Último mes cerrado con datos reales (n° de mes)", D.get("corte"), NF_N,
         "Los meses hasta este número usan solo la hoja Reales y no cambian con los escenarios. Vacío = sin datos reales."),
        ("n_InicioReal", "Mes real de inicio de operación (n° de mes)", D.get("inicio_real"), NF_N,
         "Cargar cuando la operación ya comenzó. Prevalece sobre el supuesto vigente y el atraso del escenario."),
        ("n_Escenario", "Escenario para la proyección", D.get("escenario") or "Vigente", None,
         "Vigente (sin ajustes), Conservador u Optimista. Solo afecta los meses posteriores al corte."),
    ]
    for name, lab, val, fmt, ins in rows:
        put(ws, r, 2, lab)
        inp(ws, r, 3, val, fmt)
        put(ws, r, 4, ins, font=F_S)
        nm(name, "Inicio", f"C{r}")
        if name == "n_Etapa":
            dv(ws, f"C{r}", ["Implantación", "Operación"])
        if name == "n_Escenario":
            dv(ws, f"C{r}", ["Vigente", "Conservador", "Optimista"])
        r += 1
    put(ws, r, 2, "Moneda de presentación")
    put(ws, r, 3, "USD (costos en Gs convertidos con el TC de cada año)")
    r += 2
    section(ws, r, "2. Estado del modelo (automático)", 4); r += 1
    est = [
        ("Fecha del último mes cerrado", '=IF(Corte=0,"Sin datos reales",IF(n_Fecha1="","Mes "&Corte,"Mes "&Corte&" ("&TEXT(MONTH(EOMONTH(n_Fecha1,Corte-1)),"00")&"/"&YEAR(EOMONTH(n_Fecha1,Corte-1))&")"))'),
        ("Inicio de operación (presupuesto / vigente)", '="Mes "&P_Inicio&" / mes "&V_Inicio'),
        ("Fin del horizonte evaluado (fijo)", '="Mes "&Fin&" ("&P_Anios&" años de operación desde el inicio presupuestado)"'),
        ("Supuestos vigentes modificados respecto del presupuesto", "=c_Modificados"),
        ("Controles con alerta (hoja Control)", "=c_Alertas"),
        ("Brecha de caja máxima vigente (USD)", "=r_BrechaV"),
    ]
    for lab, f in est:
        put(ws, r, 2, lab)
        put(ws, r, 3, f, font=F_B, fmt=NF_USD)
        r += 1
    r += 1
    section(ws, r, "3. Cómo cargar y actualizar", 4); r += 1
    pasos = [
        ("Alta de un negocio", "Copiar el archivo vacío, completar la sección 1 y luego las hojas en este orden: Supuestos → Productos → Costos → Inversion."),
        ("Presupuesto original", "Cargar en las columnas «Presupuesto original» los datos del plan aprobado. No se modifican después: son la referencia de comparación."),
        ("Proyección vigente", "Las columnas «Vigente» heredan el presupuesto (gris). Sobrescribir solo lo que cambió; la celda pasa a azul/amarillo y Control lo cuenta."),
        ("Cierre mensual", "1) Cargar en Reales el mes cerrado (unidades, ventas, costos, CAPEX, financiamiento). 2) Cargar los saldos al cierre en el mes de corte. "
                           "3) Actualizar «Último mes cerrado» arriba. 4) Revisar Control. 5) Ajustar la proyección vigente si cambió."),
        ("Inversión", "En Inversion registrar por ítem lo ejecutado, pagado, comprometido y por contratar. El costo final estimado se calcula solo."),
        ("Escenarios", "Elegir el escenario arriba. Los ajustes están en Supuestos sección C y solo afectan la proyección (nunca los meses reales)."),
        ("Lectura", "Resultados compara presupuesto original vs vigente (real + proyección). Control muestra los faltantes y las inconsistencias."),
    ]
    for a, b in pasos:
        put(ws, r, 2, a, font=F_B)
        put(ws, r, 3, b, al=WRAP)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 26
        r += 1
    r += 1
    section(ws, r, "4. Convenciones", 4); r += 1
    conv = [(FL_IN, F_IN, "Celda para completar (dato de entrada)."),
            (FL_INH, F_INH, "Vigente heredado del presupuesto (fórmula). Sobrescribir si el dato cambió."),
            (None, F, "Fórmula: no editar."),
            (FL_P, F, "Presupuesto original (Resultados)."),
            (FL_V, F, "Vigente: real + proyección (Resultados).")]
    for fill, font, txt in conv:
        put(ws, r, 2, "Ejemplo 1.234", font=font, fill=fill)
        put(ws, r, 3, txt)
        r += 1
    put(ws, r, 2, "Porcentajes como fracción (0,05 = 5%). Montos en USD salvo que la columna Moneda indique PYG.", font=F_S)
    r += 2
    section(ws, r, "5. Hojas", 4); r += 1
    hojas = [("Supuestos", "Parámetros generales, supuestos por año, escenarios y referencia del plan publicado."),
             ("Productos", f"Hasta {NP} productos: precio, costos unitarios, capacidad y volumen anual (presupuesto y vigente)."),
             ("Costos", "Personal por puesto (dotación por año) y costos fijos (moneda, crecimiento, inicio)."),
             ("Inversion", "CAPEX por ítem (presupuesto, ejecutado, comprometido, por contratar), préstamos y aportes."),
             ("Reales", "Carga mensual de la ejecución y saldos al cierre (formato fijo: un dato por celda)."),
             ("Calculo", "Motor mensual automático (no editar). Bloque P = presupuesto; bloque V = vigente."),
             ("Resultados", "Indicadores, resúmenes anuales, desvíos y gráficos."),
             ("Control", "Validaciones y datos faltantes.")]
    for a, b in hojas:
        put(ws, r, 2, a, font=F_B)
        put(ws, r, 3, b)
        r += 1
    r += 1
    section(ws, r, "6. Datos pendientes de este negocio (editable)", 4); r += 1
    pend = D.get("pendientes", [])
    for i in range(10):
        put(ws, r, 2, f"{i + 1}.")
        inp(ws, r, 3, pend[i] if i < len(pend) else None)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        r += 1
    ws.freeze_panes = "A4"


# ============================================================================ SUPUESTOS
PARAMS = [
    # clave, etiqueta, unidad, formato, aplica escenario
    ("Inicio", "Mes de inicio de la operación (n° de mes del modelo)", "mes", NF_N),
    ("Anios", "Años de operación evaluados", "años", NF_N),
    ("Merma", "Merma entre producción y venta", "%", NF_PCT2),
    ("Cargas", "Cargas sociales sobre salarios (aguinaldo, IPS, vacaciones)", "%", NF_PCT2),
    ("Alim", "Alimentación del personal", "Gs/persona/mes", NF_INT),
    ("PreDir", "Personal directo activo antes del inicio de operación", "% dotación año 1", NF_PCT),
    ("PreInd", "Personal indirecto y administrativo activo antes del inicio", "% dotación año 1", NF_PCT),
    ("Desp", "Despacho y gastos de importación (s/ materia prima + insumos)", "%", NF_PCT2),
    ("Pkg", "Packaging (s/ materiales, otros variables y personal directo)", "%", NF_PCT2),
    ("ImpP", "Imprevistos de producción (s/ costos de producción)", "%", NF_PCT2),
    ("ImpA", "Imprevistos de administración (s/ gastos marcados como base)", "%", NF_PCT2),
    ("Trib", "Tributo de maquila (s/ ventas)", "%", NF_PCT2),
    ("SegV", "Seguro de caución y otros cargos (s/ ventas)", "%", NF_PCT2),
    ("ComV", "Comisiones y logística de venta (s/ ventas)", "%", NF_PCT2),
    ("IRE", "Impuesto a la renta (s/ resultado)", "%", NF_PCT2),
    ("DSO", "Días de cobro a clientes", "días", NF_N),
    ("DInv", "Días de inventario de materiales", "días", NF_N),
    ("DPT", "Días de producto terminado", "días", NF_N),
    ("DPOmp", "Días de pago a proveedores de materiales", "días", NF_N),
    ("DPOot", "Días de pago a otros proveedores", "días", NF_N),
    ("CajaMin", "Caja mínima operativa", "días de ventas", NF_N),
    ("IVArec", "IVA recuperable de los costos fijos gravados", "%", NF_PCT),
    ("IVAmeses", "Meses para recuperar el IVA", "meses", NF_N),
    ("IVAint", "IVA sobre intereses", "%", NF_PCT),
    ("WACC", "Tasa de descuento del proyecto (WACC, anual)", "%", NF_PCT2),
    ("Ke", "Costo del capital propio (Ke, anual)", "%", NF_PCT2),
    ("RealCT", "Recuperación del capital de trabajo al final del horizonte", "%", NF_PCT),
    ("RealAF", "Valor de realización del activo fijo al final (s/ valor contable)", "%", NF_PCT),
    ("MesDiv", "Mes de pago de dividendos del año siguiente (1-12)", "mes", NF_N),
]
ANUAL = [("TC", "Tipo de cambio (Gs/USD) — por año del modelo", "Gs/USD", NF_INT),
         ("dPrecio", "Variación de precios de venta (año de operación)", "%", NF_PCT),
         ("dMP", "Variación del costo de materia prima", "%", NF_PCT),
         ("dOV", "Variación de insumos y otros costos variables", "%", NF_PCT),
         ("dSal", "Ajuste salarial", "%", NF_PCT),
         ("Payout", "Dividendos (% del resultado del año)", "%", NF_PCT)]
IDX = [("Precio", "dPrecio"), ("MP", "dMP"), ("OV", "dOV"), ("Sal", "dSal")]
ESC = [("dVol", "Ajuste de volumen", "%", NF_PCT, [0, -0.10, 0.05]),
       ("dPrecio", "Ajuste de precio", "%", NF_PCT, [0, -0.05, 0.02]),
       ("dMat", "Ajuste de costo de materiales y variables", "%", NF_PCT, [0, 0.10, -0.03]),
       ("dTC", "Ajuste del tipo de cambio (+ = más Gs por USD)", "%", NF_PCT, [0, 0, 0]),
       ("Atraso", "Atraso adicional del inicio de operación", "meses", NF_N, [0, 3, 0]),
       ("dCapex", "Sobrecosto sobre CAPEX pendiente", "%", NF_PCT, [0, 0.10, 0]),
       ("DSOesc", "Días de cobro (vacío = vigente)", "días", NF_N, [None, 150, 90]),
       ("DInvEsc", "Días de inventario (vacío = vigente)", "días", NF_N, [None, None, 90])]
SUPR = {}


def build_supuestos():
    ws = WS["Supuestos"]
    title(ws, "Supuestos — presupuesto original y vigente",
          "Presupuesto original = plan aprobado (no se modifica). Vigente = hereda el presupuesto; sobrescribir solo lo que cambió.")
    widths(ws, {"A": 7, "B": 58, "C": 15, "D": 14, "E": 14, "F": 14, "G": 14, "H": 13, "I": 13, "J": 13, "K": 13, "L": 13, "M": 13, "N": 60})
    r = 4
    section(ws, r, "A. Parámetros generales", 14); r += 1
    header(ws, r, ["Cód.", "Parámetro", "Unidad", "Presupuesto original", "Vigente", "Nota / fuente"], 1, 30)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
    r += 1
    for i, (k, lab, unit, fmt) in enumerate(PARAMS):
        SUPR[k] = r
        put(ws, r, 1, f"G{i + 1:02d}")
        put(ws, r, 2, lab)
        put(ws, r, 3, unit)
        inp(ws, r, 4, P(k), fmt)
        note = D.get("params", {}).get(k, (None, ""))[1] or ""
        if k == "Anios":
            put(ws, r, 5, f"=D{r}", fmt=fmt)
            note = (note + " " if note else "") + "Horizonte común a ambas versiones: se define solo en el presupuesto."
        else:
            inh(ws, r, 5, f"=D{r}", fmt)
        put(ws, r, 6, note, font=F_S)
        nm(f"o_{k}", "Supuestos", f"D{r}")
        nm(f"vg_{k}", "Supuestos", f"E{r}")
        r += 1
    SUPR["p_last"] = r - 1
    r += 1
    section(ws, r, "B. Supuestos por año (año 1 = primer año de operación; el TC se aplica por año del modelo)", 14); r += 1
    header(ws, r, ["Cód.", "Concepto", "Unidad"] + [f"Año {y}" for y in range(1, NY + 1)] + ["Nota"], 1)
    r += 1
    for ver, lab_ver, fill in [("O", "PRESUPUESTO ORIGINAL", FL_P), ("V", "VIGENTE (hereda el presupuesto)", FL_V)]:
        put(ws, r, 2, lab_ver, font=F_B, fill=fill)
        r += 1
        for k, lab, unit, fmt in ANUAL:
            SUPR[f"{k}_{ver}"] = r
            put(ws, r, 1, k)
            put(ws, r, 2, lab)
            put(ws, r, 3, unit)
            vals = D.get("anual", {}).get(k, [None] * NY)
            for y in range(1, NY + 1):
                c = 3 + y
                if ver == "O":
                    inp(ws, r, c, vals[y - 1] if vals else None, fmt)
                else:
                    inh(ws, r, c, f"={CL(c)}{SUPR[k + '_O']}", fmt)
            nm(f"{k}_{ver}", "Supuestos", f"D{r}:M{r}")
            r += 1
        for k, src in IDX:
            SUPR[f"Idx{k}_{ver}"] = r
            put(ws, r, 2, f"Índice acumulado de {dict(Precio='precios', MP='materia prima', OV='otros variables', Sal='salarios')[k]} (cálculo)", font=F_S)
            for y in range(1, NY + 1):
                c = 3 + y
                f = "=1" if y == 1 else f"={CL(c - 1)}{r}*(1+N({CL(c)}{SUPR[src + '_' + ver]}))"
                put(ws, r, c, f, fmt=NF_USD4)
            nm(f"Idx{k}_{ver}", "Supuestos", f"D{r}:M{r}")
            r += 1
        r += 1
    section(ws, r, "C. Escenarios (se aplican solo a la proyección vigente, después del último mes cerrado)", 14); r += 1
    header(ws, r, ["Cód.", "Ajuste", "Unidad", "Vigente", "Conservador", "Optimista", "Aplicado", "Nota"], 1, 30)
    ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=14)
    r += 1
    SUPR["esc_first"] = r
    for i, (k, lab, unit, fmt, vals) in enumerate(ESC):
        put(ws, r, 1, f"E{i + 1}")
        put(ws, r, 2, lab)
        put(ws, r, 3, unit)
        put(ws, r, 4, vals[0], fmt=fmt, fill=FL_TOT)
        inp(ws, r, 5, vals[1], fmt)
        inp(ws, r, 6, vals[2], fmt)
        put(ws, r, 7, f'=IF(INDEX(D{r}:F{r},1,c_EscIdx)="","",INDEX(D{r}:F{r},1,c_EscIdx))', fmt=fmt, font=F_B)
        nm(f"e_{k}", "Supuestos", f"G{r}")
        r += 1
    put(ws, r, 2, "Índice del escenario activo")
    put(ws, r, 4, '=IF(n_Escenario="Conservador",2,IF(n_Escenario="Optimista",3,1))', fmt=NF_N)
    nm("c_EscIdx", "Supuestos", f"D{r}")
    put(ws, r, 8, "Conservador y Optimista son ajustes ILUSTRATIVOS editables, no premisas aprobadas.", font=F_S)
    r += 2
    section(ws, r, "D. Referencia del plan publicado (opcional; solo para comparar)", 14); r += 1
    header(ws, r, ["Cód.", "Concepto", "Unidad"] + [f"Año {y}" for y in range(1, NY + 1)] + ["Nota"], 1)
    r += 1
    pub = D.get("publicado", {})
    for k, lab in [("Ventas", "Ventas publicadas"), ("EBITDA", "EBITDA publicado"), ("Utilidad", "Resultado neto publicado")]:
        SUPR[f"pub_{k}"] = r
        put(ws, r, 1, k)
        put(ws, r, 2, lab)
        put(ws, r, 3, "USD")
        vals = pub.get(k, [None] * NY)
        for y in range(1, NY + 1):
            inp(ws, r, 3 + y, vals[y - 1] if vals else None, NF_USD)
        nm(f"pub_{k}", "Supuestos", f"D{r}:M{r}")
        r += 1
    for k, lab, fmt in [("VAN", "VAN publicado", NF_USD), ("TIR", "TIR publicada", NF_PCT2)]:
        put(ws, r, 1, k)
        put(ws, r, 2, lab)
        inp(ws, r, 4, pub.get(k), fmt)
        nm(f"pub_{k}", "Supuestos", f"D{r}")
        r += 1
    put(ws, r, 2, pub.get("nota", ""), font=F_S)
    r += 2
    section(ws, r, "E. Parámetros aplicados por versión (cálculo automático; no editar)", 14); r += 1
    header(ws, r, ["", "Parámetro", "", "Presupuesto (P)", "Vigente (V)"], 1)
    r += 1
    for k, lab, unit, fmt in PARAMS:
        put(ws, r, 2, lab, font=F_S)
        put(ws, r, 4, "=MAX(1,N(o_Inicio))" if k == "Inicio" else f"=N(o_{k})", fmt=fmt)
        if k == "DSO":
            fv = '=IF(e_DSOesc="",N(vg_DSO),e_DSOesc)'
        elif k == "DInv":
            fv = '=IF(e_DInvEsc="",N(vg_DInv),e_DInvEsc)'
        elif k == "Inicio":
            fv = '=MAX(1,IF(ISNUMBER(n_InicioReal),n_InicioReal,IF(N(vg_Inicio)<=Corte,N(vg_Inicio),N(vg_Inicio)+ROUND(N(e_Atraso),0))))'
        else:
            fv = f"=N(vg_{k})"
        put(ws, r, 5, fv, fmt=fmt)
        nm(f"P_{k}", "Supuestos", f"D{r}")
        nm(f"V_{k}", "Supuestos", f"E{r}")
        r += 1
    for k, lab in [("dVol", "Ajuste de volumen"), ("dPrecio", "Ajuste de precio"), ("dMat", "Ajuste de materiales"), ("dTC", "Ajuste de TC"),
                   ("dCapex", "Sobrecosto CAPEX pendiente")]:
        put(ws, r, 2, lab, font=F_S)
        put(ws, r, 4, 0, fmt=NF_PCT)
        put(ws, r, 5, f"=N(e_{k})", fmt=NF_PCT)
        nm(f"P_{k}", "Supuestos", f"D{r}")
        nm(f"V_{k}", "Supuestos", f"E{r}")
        r += 1
    for name, lab, f, fmt in [("Corte", "Último mes cerrado (0 = sin reales)", "=IF(ISNUMBER(n_Corte),n_Corte,0)", NF_N),
                              ("Fin", "Último mes del horizonte (fijo, según presupuesto)", "=P_Inicio+12*P_Anios-1", NF_N),
                              ("DiasMes", "Días por mes", "=365/12", NF_USD2)]:
        put(ws, r, 2, lab, font=F_S)
        put(ws, r, 4, f, fmt=fmt)
        nm(name, "Supuestos", f"D{r}")
        r += 1
    ws.freeze_panes = "D6"


# ============================================================================ PRODUCTOS
PR = {}


def build_productos():
    ws = WS["Productos"]
    title(ws, "Productos — precios, costos unitarios, capacidad y volúmenes (USD por unidad de venta)",
          f"Hasta {NP} productos. Las columnas «vigente» heredan el presupuesto; sobrescribir lo que cambió.")
    widths(ws, {"A": 4, "B": 34, "C": 8})
    for c in range(4, 16):
        ws.column_dimensions[CL(c)].width = 12.5
    ws.column_dimensions["O"].width = 50
    r = 4
    section(ws, r, "A. Datos por producto", 15); r += 1
    header(ws, r, ["#", "Producto", "Unidad", "Capacidad inicial (u/año)", "Capacidad ampliada (u/año)", "Desde año de operación",
                   "Precio año 1 — presupuesto", "Precio año 1 — vigente", "Materia prima /u — presupuesto", "Materia prima /u — vigente",
                   "Insumos /u — presupuesto", "Insumos /u — vigente", "Otros variables /u — presupuesto", "Otros variables /u — vigente", "Nota"], 1, 54)
    r += 1
    PR["first"] = r
    prods = D.get("productos", [])
    for i in range(NP):
        p = prods[i] if i < len(prods) else None
        put(ws, r, 1, i + 1)
        inp(ws, r, 2, p[0] if p else None)
        inp(ws, r, 3, p[1] if p else None)
        inp(ws, r, 4, p[2] if p else None, NF_INT)
        inp(ws, r, 5, p[3] if p else None, NF_INT)
        inp(ws, r, 6, p[4] if p else None, NF_N)
        for j, (co, cv) in enumerate([(7, 8), (9, 10), (11, 12), (13, 14)]):
            inp(ws, r, co, p[5 + j] if p else None, NF_USD4)
            inh(ws, r, cv, f"={CL(co)}{r}", NF_USD4)
        put(ws, r, 15, p[9] if p else None, font=F_S)
        r += 1
    PR["last"] = r - 1
    f, l = PR["first"], PR["last"]
    for name, col in [("Prod_Nombre", "B"), ("Cap_Ini", "D"), ("Cap_Amp", "E"), ("Cap_Anio", "F"), ("Precio_O", "G"), ("Precio_V", "H"),
                      ("MP_O", "I"), ("MP_V", "J"), ("INS_O", "K"), ("INS_V", "L"), ("OV_O", "M"), ("OV_V", "N")]:
        nm(name, "Productos", f"{col}{f}:{col}{l}")
    r += 1
    vols = D.get("volumen", [])
    for ver, lab, fill in [("O", "B. Ventas en unidades por año de operación — PRESUPUESTO ORIGINAL", FL_P),
                           ("V", "C. Ventas en unidades por año de operación — VIGENTE (hereda; sobrescribir lo que cambió)", FL_V)]:
        section(ws, r, lab, 15); r += 1
        header(ws, r, ["#", "Producto", "Unidad"] + [f"Año {y}" for y in range(1, NY + 1)] + ["Total"], 1)
        r += 1
        PR[f"vol_{ver}"] = r
        for i in range(NP):
            put(ws, r, 1, i + 1)
            put(ws, r, 2, f'=IF({CL(2)}{f + i}="","",{CL(2)}{f + i})', font=F_LINK)
            put(ws, r, 3, f'=IF({CL(3)}{f + i}="","",{CL(3)}{f + i})', font=F_LINK)
            for y in range(1, NY + 1):
                c = 3 + y
                if ver == "O":
                    v = vols[i][y - 1] if i < len(vols) else None
                    inp(ws, r, c, v, NF_INT)
                else:
                    inh(ws, r, c, f"={CL(c)}{PR['vol_O'] + i}", NF_INT)
            put(ws, r, 14, f"=SUM(D{r}:M{r})", fmt=NF_INT, font=F_B)
            r += 1
        nm(f"VolTot_{ver}", "Productos", f"N{PR['vol_' + ver]}:N{PR['vol_' + ver] + NP - 1}")
        put(ws, r, 2, "Total", font=F_B)
        for y in range(1, NY + 2):
            c = 3 + y
            put(ws, r, c, f"=SUM({CL(c)}{r - NP}:{CL(c)}{r - 1})", fmt=NF_INT, font=F_B, fill=FL_TOT)
        nm(f"Vol_{ver}", "Productos", f"D{PR['vol_' + ver]}:M{PR['vol_' + ver] + NP - 1}")
        r += 2
    section(ws, r, "D. Control de capacidad — unidades vigentes por encima de la capacidad (incluye merma)", 15); r += 1
    header(ws, r, ["#", "Producto", ""] + [f"Año {y}" for y in range(1, NY + 1)] + ["Años con exceso"], 1)
    r += 1
    PR["cap_first"] = r
    for i in range(NP):
        put(ws, r, 1, i + 1)
        put(ws, r, 2, f'=IF(B{f + i}="","",B{f + i})', font=F_LINK)
        for y in range(1, NY + 1):
            c = 3 + y
            vol = f"{CL(c)}{PR['vol_V'] + i}"
            cap = f"IF(AND(N($F{f + i})>0,{y}>=$F{f + i}),N($E{f + i}),N($D{f + i}))"
            put(ws, r, c, f"=IF({cap}=0,0,MAX(0,{vol}/(1-V_Merma)-{cap}))", fmt=NF_INT)
        put(ws, r, 14, f'=SUMPRODUCT(--(D{r}:M{r}>0.5))', fmt=NF_N, font=F_B)
        r += 1
    nm("c_CapExceso", "Productos", f"N{PR['cap_first']}:N{r - 1}")
    put(ws, r, 2, "Capacidad 0 = no informada (no se controla).", font=F_S)
    ws.freeze_panes = "C6"


# ============================================================================ COSTOS
CO = {}


def build_costos():
    ws = WS["Costos"]
    title(ws, "Costos — personal por puesto y costos fijos",
          "Salarios y costos fijos pueden cargarse en PYG o USD (columna Moneda). Se convierten con el TC de cada año.")
    widths(ws, {"A": 4, "B": 40, "C": 16, "D": 9, "E": 14, "F": 14})
    for c in range(7, 17):
        ws.column_dimensions[CL(c)].width = 10
    ws.column_dimensions["Q"].width = 40
    r = 4
    section(ws, r, "A. Personal por puesto — salario mensual y dotación por año de operación (PRESUPUESTO ORIGINAL; salario vigente al lado)", 17); r += 1
    header(ws, r, ["#", "Puesto", "Clase", "Moneda", "Salario mensual — presupuesto", "Salario mensual — vigente"] +
           [f"Dot. año {y}" for y in range(1, NY + 1)] + ["Nota"], 1, 40)
    r += 1
    CO["per_first"] = r
    pers = D.get("personal", [])
    for i in range(NPER):
        p = pers[i] if i < len(pers) else None
        put(ws, r, 1, i + 1)
        inp(ws, r, 2, p[0] if p else None)
        inp(ws, r, 3, p[1] if p else None)
        inp(ws, r, 4, p[2] if p else None)
        inp(ws, r, 5, p[3] if p else None, NF_INT)
        inh(ws, r, 6, f"=E{r}", NF_INT)
        for y in range(1, NY + 1):
            v = (p[4][min(y, len(p[4])) - 1]) if p else None
            inp(ws, r, 6 + y, v, NF_INT)
        r += 1
    CO["per_last"] = r - 1
    f, l = CO["per_first"], CO["per_last"]
    dv(ws, f"C{f}:C{l}", ["Directo", "Indirecto", "Administración"])
    dv(ws, f"D{f}:D{l}", ["PYG", "USD"])
    nm("Per_Clase", "Costos", f"C{f}:C{l}")
    nm("Per_Mon", "Costos", f"D{f}:D{l}")
    nm("Sal_O", "Costos", f"E{f}:E{l}")
    nm("Sal_V", "Costos", f"F{f}:F{l}")
    nm("Dot_O", "Costos", f"G{f}:P{l}")
    put(ws, r, 2, "Totales de dotación", font=F_B)
    for y in range(1, NY + 1):
        put(ws, r, 6 + y, f"=SUM({CL(6 + y)}{f}:{CL(6 + y)}{l})", fmt=NF_INT, font=F_B, fill=FL_TOT)
    r += 1
    put(ws, r, 2, "Clase: Directo = mano de obra de producción; Indirecto = supervisión, calidad, mantenimiento, depósito (costo de producción); "
                  "Administración = gerencia general y administración.", font=F_S)
    r += 2
    section(ws, r, "B. Dotación VIGENTE por año de operación (hereda; sobrescribir lo que cambió)", 17); r += 1
    header(ws, r, ["#", "Puesto", "Clase", "", "", ""] + [f"Dot. año {y}" for y in range(1, NY + 1)], 1)
    r += 1
    CO["dotv_first"] = r
    for i in range(NPER):
        put(ws, r, 1, i + 1)
        put(ws, r, 2, f'=IF(B{f + i}="","",B{f + i})', font=F_LINK)
        put(ws, r, 3, f'=IF(C{f + i}="","",C{f + i})', font=F_LINK)
        for y in range(1, NY + 1):
            inh(ws, r, 6 + y, f"={CL(6 + y)}{f + i}", NF_INT)
        r += 1
    nm("Dot_V", "Costos", f"G{CO['dotv_first']}:P{r - 1}")
    put(ws, r, 2, "Totales de dotación vigente", font=F_B)
    for y in range(1, NY + 1):
        put(ws, r, 6 + y, f"=SUM({CL(6 + y)}{CO['dotv_first']}:{CL(6 + y)}{r - 1})", fmt=NF_INT, font=F_B, fill=FL_TOT)
    r += 2
    section(ws, r, "C. Costos fijos y gastos (monto mensual del año 1)", 17); r += 1
    header(ws, r, ["#", "Concepto", "Clase", "Moneda", "Monto mensual — presupuesto", "Monto mensual — vigente", "Crecimiento anual",
                   "Desde año de operación (0 = implantación)", "IVA incluido", "Base de imprevistos (1/0)", "Nota"], 1, 54)
    ws.merge_cells(start_row=r, start_column=11, end_row=r, end_column=17)
    r += 1
    CO["fc_first"] = r
    fij = D.get("fijos", [])
    for i in range(NFIJ):
        p = fij[i] if i < len(fij) else None
        put(ws, r, 1, i + 1)
        inp(ws, r, 2, p[0] if p else None)
        inp(ws, r, 3, p[1] if p else None)
        inp(ws, r, 4, p[2] if p else None)
        inp(ws, r, 5, p[3] if p else None, NF_USD2)
        inh(ws, r, 6, f"=E{r}", NF_USD2)
        inp(ws, r, 7, p[4] if p else None, NF_PCT)
        inp(ws, r, 8, p[5] if p else None, NF_N)
        inp(ws, r, 9, p[6] if p else None, NF_PCT)
        inp(ws, r, 10, p[7] if p else None, NF_N)
        put(ws, r, 11, p[8] if p else None, font=F_S)
        r += 1
    f2, l2 = CO["fc_first"], r - 1
    dv(ws, f"C{f2}:C{l2}", ["Producción", "Administración"])
    dv(ws, f"D{f2}:D{l2}", ["PYG", "USD"])
    for name, col in [("FC_Clase", "C"), ("FC_Mon", "D"), ("FC_O", "E"), ("FC_V", "F"), ("FC_g", "G"), ("FC_Desde", "H"), ("FC_IVA", "I"), ("FC_Base", "J")]:
        nm(name, "Costos", f"{col}{f2}:{col}{l2}")
    put(ws, r, 2, "Montos con IVA incluido cuando no es recuperable (criterio prudencial). Los imprevistos de producción se calculan sobre todos los costos de producción.",
        font=F_S)
    ws.freeze_panes = "C6"


# ============================================================================ INVERSION
IV = {}


def build_inversion():
    ws = WS["Inversion"]
    title(ws, "Inversión y financiamiento — CAPEX por ítem, préstamos y aportes (USD)",
          "Costo final estimado = ejecutado + comprometido pendiente + por contratar. Los pendientes se pagan en el «mes de pago pendiente».")
    widths(ws, {"A": 4, "B": 40, "C": 14})
    for c in range(4, 22):
        ws.column_dimensions[CL(c)].width = 12.5
    r = 4
    section(ws, r, "A. Inversiones (CAPEX)", 21); r += 1
    header(ws, r, ["#", "Ítem", "Clase", "Presupuesto original", "Mes de pago (presup.)", "Vida útil (años)",
                   "Ejecutado (devengado real)", "Pagado (real)", "Comprometido pendiente", "Por contratar (si difiere)", "Mes de pago pendiente (si difiere)",
                   "Mes de puesta en servicio real", "Por contratar usado", "Costo final estimado", "Desvío vs presupuesto", "Avance financiero",
                   "Mes de pago pendiente usado", "Pendiente devengado + comprometido", "Servicio (presup.)", "Servicio (vigente)",
                   "Base depreciable vigente"], 1, 66)
    r += 1
    IV["first"] = r
    cap = D.get("capex", [])
    for i in range(NCAP):
        p = cap[i] if i < len(cap) else None
        put(ws, r, 1, i + 1)
        inp(ws, r, 2, p[0] if p else None)
        inp(ws, r, 3, p[1] if p else None)
        inp(ws, r, 4, p[2] if p else None, NF_USD)
        inp(ws, r, 5, p[3] if p else None, NF_N)
        inp(ws, r, 6, p[4] if p else None, NF_N)
        for c in range(7, 13):
            inp(ws, r, c, None, NF_USD if c < 11 else NF_N)
        put(ws, r, 13, f'=IF(J{r}<>"",N(J{r}),MAX(0,N(D{r})-N(G{r})-N(I{r})))', fmt=NF_USD)
        put(ws, r, 14, f"=N(G{r})+N(I{r})+M{r}", fmt=NF_USD, font=F_B)
        put(ws, r, 15, f"=N{r}-N(D{r})", fmt=NF_USD)
        put(ws, r, 16, f'=IF(N{r}>0,N(H{r})/N{r},"")', fmt=NF_PCT)
        put(ws, r, 17, f'=IF(K{r}<>"",K{r},MAX(N(E{r}),Corte+1))', fmt=NF_N)
        put(ws, r, 18, f"=N(G{r})-N(H{r})+N(I{r})", fmt=NF_USD)
        put(ws, r, 19, f"=MAX(N(E{r}),P_Inicio)", fmt=NF_N)
        put(ws, r, 20, f'=IF(L{r}<>"",L{r},MAX(N(E{r}),V_Inicio))', fmt=NF_N)
        put(ws, r, 21, f"=N(G{r})+N(I{r})+M{r}*(1+V_dCapex)", fmt=NF_USD)
        r += 1
    IV["last"] = r - 1
    f, l = IV["first"], IV["last"]
    for name, col in [("Inv_D", "D"), ("Inv_E", "E"), ("Inv_F", "F"), ("Inv_H", "H"), ("Inv_M", "M"), ("Inv_Q", "Q"), ("Inv_R", "R"),
                      ("Inv_SvcP", "S"), ("Inv_SvcV", "T"), ("Inv_BaseV", "U"), ("Inv_CF", "N")]:
        nm(name, "Inversion", f"{col}{f}:{col}{l}")
    put(ws, r, 2, "TOTAL", font=F_B)
    for c in [4, 7, 8, 9, 13, 14, 15, 18, 21]:
        put(ws, r, c, f"=SUM({CL(c)}{f}:{CL(c)}{l})", fmt=NF_USD, font=F_B, fill=FL_TOT)
    put(ws, r, 16, f'=IF(N{r}>0,H{r}/N{r},"")', fmt=NF_PCT, font=F_B)
    IV["tot"] = r
    nm("i_CapexP", "Inversion", f"D{r}")
    nm("i_CapexV", "Inversion", f"N{r}")
    nm("i_Pagado", "Inversion", f"H{r}")
    r += 1
    put(ws, r, 2, D.get("capex_nota") or "Aportes en especie (máquinas): cargar el ítem aquí y el mismo monto como aporte en el mismo mes.", font=F_S)
    r += 2
    section(ws, r, "B. Préstamos (sistema francés; desembolso al inicio del mes; tasa mensual = nominal anual / 12)", 21); r += 1
    header(ws, r, ["", "Condición", "Unidad", "Préstamo 1 — presupuesto", "Préstamo 1 — vigente", "Préstamo 2 — presupuesto", "Préstamo 2 — vigente",
                   "Préstamo 3 — presupuesto", "Préstamo 3 — vigente"], 1, 42)
    r += 1
    loans = D.get("prestamos", [])
    for k, lab, unit, fmt, idx in [("Nombre", "Nombre / banco", "", None, 0), ("Monto", "Monto", "USD", NF_USD, 1),
                                   ("Mes", "Mes de desembolso", "mes", NF_N, 2), ("Tasa", "Tasa nominal anual", "%", NF_PCT2, 3),
                                   ("Gracia", "Meses de gracia (solo intereses)", "meses", NF_N, 4), ("Cuotas", "Cuotas de capital e interés", "cuotas", NF_N, 5)]:
        put(ws, r, 2, lab)
        put(ws, r, 3, unit)
        for j in range(NL):
            lo = loans[j] if j < len(loans) else None
            co, cv = 4 + 2 * j, 5 + 2 * j
            inp(ws, r, co, lo[idx] if lo else None, fmt)
            inh(ws, r, cv, f"={CL(co)}{r}", fmt)
            if k != "Nombre":
                nm(f"P_L{j + 1}_{k}", "Inversion", f"{CL(co)}{r}")
                nm(f"V_L{j + 1}_{k}", "Inversion", f"{CL(cv)}{r}")
        r += 1
    put(ws, r, 2, "Préstamo de corto plazo con cronograma (p. ej. 12 cuotas o pago único): cargarlo como préstamo 2 o 3.", font=F_S)
    r += 2
    section(ws, r, "B2. Línea de crédito rotativa de corto plazo (capital de trabajo): se usa solo si la caja cae bajo el mínimo y se devuelve con el excedente", 21); r += 1
    header(ws, r, ["", "Condición", "Unidad", "Presupuesto", "Vigente", "Nota"], 1, 30)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
    r += 1
    lc = D.get("linea", {})
    for k, lab, unit, fmt, note in [("Lim", "Límite aprobado", "USD", NF_USD, "0 o vacío = sin línea."),
                                    ("Tasa", "Tasa nominal anual (interés mensual sobre el saldo usado)", "%", NF_PCT2, "Más IVA sobre intereses (Supuestos G24)."),
                                    ("Desde", "Disponible desde el mes", "mes", NF_N, "Vacío = desde el mes 1."),
                                    ("Hasta", "Vencimiento: último mes de uso (luego se cancela el saldo)", "mes", NF_N,
                                     "Vacío = fin del horizonte. Al vencer se paga todo el saldo aunque falte caja (riesgo de renovación).")]:
        put(ws, r, 2, lab)
        put(ws, r, 3, unit)
        inp(ws, r, 4, lc.get(k), fmt)
        inh(ws, r, 5, f"=D{r}", fmt)
        put(ws, r, 6, note, font=F_S)
        nm(f"P_LC_{k}", "Inversion", f"D{r}")
        nm(f"V_LC_{k}", "Inversion", f"E{r}")
        r += 1
    put(ws, r, 2, "Saldo utilizado de la línea al último mes cerrado (solo vigente)")
    put(ws, r, 3, "USD")
    inp(ws, r, 5, None, NF_USD)
    put(ws, r, 6, "Los usos y devoluciones reales se cargan en Reales (desembolsos / amortización); aquí solo el saldo al corte.", font=F_S)
    nm("V_LC_Saldo0", "Inversion", f"E{r}")
    r += 1
    put(ws, r, 2, "Los dividendos solo se pagan con caja sobre el mínimo después de devolver la línea.", font=F_S)
    r += 2
    section(ws, r, "C. Aportes de socios (efectivo o especie)", 21); r += 1
    header(ws, r, ["#", "Concepto", "Mes — presupuesto", "Monto — presupuesto", "Mes — vigente", "Monto — vigente", "Nota"], 1, 30)
    r += 1
    IV["ap_first"] = r
    aps = D.get("aportes", [])
    for i in range(NAP):
        a = aps[i] if i < len(aps) else None
        put(ws, r, 1, i + 1)
        inp(ws, r, 2, a[0] if a else None)
        inp(ws, r, 3, a[1] if a else None, NF_N)
        inp(ws, r, 4, a[2] if a else None, NF_USD)
        inh(ws, r, 5, f"=C{r}", NF_N)
        inh(ws, r, 6, f"=D{r}", NF_USD)
        r += 1
    fa, la = IV["ap_first"], r - 1
    nm("Ap_MesO", "Inversion", f"C{fa}:C{la}")
    nm("Ap_MontoO", "Inversion", f"D{fa}:D{la}")
    nm("Ap_MesV", "Inversion", f"E{fa}:E{la}")
    nm("Ap_MontoV", "Inversion", f"F{fa}:F{la}")
    ws.freeze_panes = "C6"


# ============================================================================ REALES
RR = {}
REAL_ROWS = [
    ("sec", "Tipo de cambio"),
    ("TC", "Tipo de cambio real del mes (Gs/USD)", "Gs/USD"),
    ("sec", "Unidades vendidas por producto"),
] + [(f"U{i + 1}", None, "u") for i in range(NP)] + [
    ("sec", "Ventas por producto (facturación neta)"),
] + [(f"S{i + 1}", None, "mon") for i in range(NP)] + [
    ("sec", "Costos y gastos devengados (sin IVA recuperable)"),
    ("MAT", "Materia prima e insumos consumidos", "mon"),
    ("OVAR", "Otros costos variables (flete, despacho, packaging, seguros, comisiones)", "mon"),
    ("PERP", "Personal de producción (directo e indirecto, con cargas)", "mon"),
    ("PERA", "Personal administrativo (con cargas)", "mon"),
    ("FIJP", "Costos fijos de producción (energía, mantenimiento, otros)", "mon"),
    ("ADM", "Gastos de administración (alquiler, servicios, honorarios, otros)", "mon"),
    ("TRIB", "Tributo de maquila", "mon"),
    ("sec", "Inversión y financiamiento (movimientos de caja)"),
    ("CAPEX", "CAPEX pagado", "mon"),
    ("DESEMB", "Desembolsos de préstamos recibidos (incluye línea rotativa)", "mon"),
    ("INT", "Intereses, IVA y comisiones pagados", "mon"),
    ("AMORT", "Amortización de préstamos pagada (incluye devoluciones de la línea)", "mon"),
    ("APORTE", "Aportes de socios recibidos", "mon"),
    ("DIV", "Dividendos pagados", "mon"),
    ("sec", "Saldos al cierre del mes (obligatorio en el último mes cerrado)"),
    ("S_CAJA", "Caja y bancos", "mon"),
    ("S_CXC", "Cuentas por cobrar a clientes", "mon"),
    ("S_INV", "Inventarios (materiales y producto terminado)", "mon"),
    ("S_IVA", "Crédito fiscal IVA por recuperar", "mon"),
    ("S_PROV", "Proveedores y cuentas por pagar", "mon"),
    ("S_DEUDA", "Deuda financiera (control)", "mon"),
]


def build_reales():
    ws = WS["Reales"]
    title(ws, "Reales — carga mensual de la ejecución (un dato por celda; solo meses cerrados)",
          "Montos en la moneda de la columna C (USD o PYG). Un mes cerrado sin datos se informa en Control (un faltante no equivale a cero).")
    widths(ws, {"A": 9, "B": 52, "C": 8, "D": 14, "E": 4})
    for m in range(1, N + 1):
        ws.column_dimensions[mc(m)].width = 11
    r = 4
    labels = [("Mes n°", None), ("Fecha", NF_MES), ("Año del modelo", NF_N), ("Estado", None)]
    for i, (lab, fmt) in enumerate(labels):
        put(ws, r + i, 2, lab, font=F_H, fill=FL_H)
        for c in (1, 3, 4, 5):
            ws.cell(row=r + i, column=c).fill = FL_H
        for m in range(1, N + 1):
            if i == 0:
                v = m
            elif i == 1:
                v = f'=IF(n_Fecha1="","",EOMONTH(n_Fecha1,{m}-1))'
            elif i == 2:
                v = f"=ROUNDUP({m}/12,0)"
            else:
                v = f'=IF({m}<=Corte,"Cerrado","Abierto")'
            put(ws, r + i, FCOL + m - 1, v, font=F_H, fill=FL_H, fmt=fmt, al=CEN)
    put(ws, r, 4, "Total cargado", font=F_H, fill=FL_H)
    RR["hdr"] = r
    r += 5
    for it in REAL_ROWS:
        if it[0] == "sec":
            section(ws, r, it[1], 5)
            r += 1
            continue
        code, lab, unit = it
        RR[code] = r
        put(ws, r, 1, code, font=F_S)
        if code.startswith("U"):
            i = int(code[1:]) - 1
            lab = f'=IF(INDEX(Prod_Nombre,{i + 1})="","Producto {i + 1} (unidades)",INDEX(Prod_Nombre,{i + 1})&" (unidades)")'
            put(ws, r, 2, lab, font=F_LINK)
            put(ws, r, 3, "u")
        elif code.startswith("S") and code[1:].isdigit():
            i = int(code[1:]) - 1
            lab = f'=IF(INDEX(Prod_Nombre,{i + 1})="","Producto {i + 1} (ventas)",INDEX(Prod_Nombre,{i + 1})&" (ventas)")'
            put(ws, r, 2, lab, font=F_LINK)
            inp(ws, r, 3, "USD")
        else:
            put(ws, r, 2, lab)
            if unit == "mon":
                inp(ws, r, 3, "PYG" if code in ("PERP", "PERA") else "USD")
            else:
                put(ws, r, 3, unit)
        put(ws, r, 4, f"=SUM({M1}{r}:{MN}{r})" if not code.startswith("S_") and code != "TC" else None, fmt=NF_USD, font=F_B)
        for m in range(1, N + 1):
            inp(ws, r, FCOL + m - 1, None, NF_USD)
        r += 1
    dv(ws, f"C{RR['S1']}:C{RR['S_DEUDA']}", ["USD", "PYG"])
    ws.conditional_formatting.add(f"{M1}{RR['hdr'] + 3}:{MN}{RR['hdr'] + 3}",
                                  CellIsRule(operator="equal", formula=['"Cerrado"'], fill=FL_V, font=Font(name=FONT, size=9, bold=True, color="000000")))
    ws.freeze_panes = f"{M1}{RR['hdr'] + 4}"


# ============================================================================ CALCULO
ROW = {}
CUR = ["Calculo"]


def A(key, m):
    return f"{mc(m)}${ROW[key]}"


def RG(key, sheet=None):
    pre = f"{sheet}!" if sheet else ""
    return f"{pre}${M1}${ROW[key]}:${MN}${ROW[key]}"


def prv(key, m, d="0"):
    return A(key, m - 1) if m > 1 else d


def rv(code, m):
    """Valor real convertido a USD (Reales), con el TC real del mes o, si falta, el vigente."""
    c, r = mc(m), RR[code]
    if code.startswith("U"):
        return f"N(Reales!{c}{r})"
    return f'IF(Reales!$C${r}="PYG",IF({A("tcr", m)}>0,N(Reales!{c}{r})/{A("tcr", m)},0),N(Reales!{c}{r}))'


def rgiven(code, m):
    return f"ISNUMBER(Reales!{mc(m)}{RR[code]})"


def version_rows(X):
    """Filas del motor para la versión X ('P' presupuesto, 'V' vigente). Mismas fórmulas."""
    O = "O" if X == "P" else "V"
    real = X == "V"
    K = lambda k: f"{X}.{k}"
    a = lambda k, m: A(K(k), m)
    hor = lambda m: A("hor", m)

    def R(code, m, fc):
        return f"=IF({A('real', m)}=1,{rv(code, m)},{fc})" if real else f"={fc}"

    rows = []
    rows.append(("sec", "Calendario de la versión"))
    rows += [
        ("aop", "Año de operación (0 = fuera de operación)", "n°", lambda m: f"=IF(AND({m}>={X}_Inicio,{m}<=Fin),INT(({m}-{X}_Inicio)/12)+1,0)", NF_N, None),
        ("op", "Operación (1/0)", "1/0", lambda m: f"=IF({a('aop', m)}>0,1,0)", NF_N, "sum"),
        ("pre", "Implantación: antes del inicio de operación (1/0)", "1/0", lambda m: f"=IF(AND({m}<{X}_Inicio,{m}<=Fin),1,0)", NF_N, "sum"),
        ("ya", "Índice de año para crecimientos (mín. 1)", "n°", lambda m: f"=MAX(1,{a('aop', m)})", NF_N, None),
        ("tc", "Tipo de cambio aplicado (Gs/USD)", "Gs/USD", lambda m: f"=INDEX(TC_{O},1,MIN({NY},{A('anio', m)}))*(1+{X}_dTC)", NF_INT, None),
        ("itc", "Inverso del TC (0 si falta el TC)", "", lambda m: f"=IF({a('tc', m)}>0,1/{a('tc', m)},0)", NF_USD4, None),
    ]
    rows.append(("sec", "Ventas"))
    for i in range(NP):
        rows.append((f"u{i + 1}", f'=IF(INDEX(Prod_Nombre,{i + 1})="","Producto {i + 1}",INDEX(Prod_Nombre,{i + 1}))&" — unidades"', "u",
                     (lambda m, i=i: (f"=IF({A('real', m)}=1,{rv('U' + str(i + 1), m)},IF({a('op', m)}=1,INDEX(Vol_{O},{i + 1},{a('aop', m)})/12*(1+{X}_dVol),0))"
                                      if real else f"=IF({a('op', m)}=1,INDEX(Vol_{O},{i + 1},{a('aop', m)})/12*(1+{X}_dVol),0)")), NF_INT, "sum"))
    for i in range(NP):
        fc = lambda m, i=i: f"{a('u' + str(i + 1), m)}*INDEX(Precio_{O},{i + 1})*INDEX(IdxPrecio_{O},1,MIN({NY},{a('ya', m)}))*(1+{X}_dPrecio)"
        rows.append((f"s{i + 1}", f'=IF(INDEX(Prod_Nombre,{i + 1})="","Producto {i + 1}",INDEX(Prod_Nombre,{i + 1}))&" — ventas USD"', "USD",
                     (lambda m, i=i, fc=fc: R("S" + str(i + 1), m, fc(m))), NF_USD, "sum"))
    rows.append(("vtas", "VENTAS TOTALES", "USD", lambda m: f"=SUM({a('s1', m)}:{a('s' + str(NP), m)})", NF_USD, "sum", True))
    urng = lambda m: f"{mc(m)}{ROW[K('u1')]}:{mc(m)}{ROW[K('u' + str(NP))]}"
    yi = lambda nmx, m: f"INDEX({nmx}_{O},1,MIN({NY},{a('ya', m)}))"
    rows.append(("sec", "Costos de la proyección (en meses reales se reemplazan por Reales)"))
    rows += [
        ("mp_f", "Materia prima", "USD", lambda m: f"=SUMPRODUCT({urng(m)},MP_{O})*{yi('IdxMP', m)}*(1+{X}_dMat)/(1-{X}_Merma)", NF_USD, "sum"),
        ("ins_f", "Insumos", "USD", lambda m: f"=SUMPRODUCT({urng(m)},INS_{O})*{yi('IdxOV', m)}*(1+{X}_dMat)/(1-{X}_Merma)", NF_USD, "sum"),
        ("ov_f", "Otros costos variables por unidad", "USD", lambda m: f"=SUMPRODUCT({urng(m)},OV_{O})*{yi('IdxOV', m)}*(1+{X}_dMat)/(1-{X}_Merma)", NF_USD, "sum"),
        ("desp_f", "Despacho de importación", "USD", lambda m: f"={X}_Desp*({a('mp_f', m)}+{a('ins_f', m)})", NF_USD, "sum"),
    ]
    act = lambda m, coef: f"({a('op', m)}+{a('pre', m)}*{coef})"
    for cls, k, coef in [("Directo", "dir", f"{X}_PreDir"), ("Indirecto", "ind", f"{X}_PreInd"), ("Administración", "adm", f"{X}_PreInd")]:
        rows.append((f"{k}_n", f"Personas activas — {cls}", "pers.",
                     (lambda m, cls=cls, coef=coef: f'=SUMPRODUCT((Per_Clase="{cls}")*INDEX(Dot_{O},0,MIN({NY},{a("ya", m)})))*{act(m, coef)}'), NF_USD2, None))
        rows.append((f"{k}_f", f"Personal — {cls} (salarios, cargas y alimentación)", "USD",
                     (lambda m, cls=cls, k=k, coef=coef: (
                         f'=SUMPRODUCT((Per_Clase="{cls}")*Sal_{O}*INDEX(Dot_{O},0,MIN({NY},{a("ya", m)}))*((Per_Mon="PYG")*{a("itc", m)}+(Per_Mon<>"PYG")))'
                         f"*{act(m, coef)}*{yi('IdxSal', m)}*(1+{X}_Cargas)+{a(k + '_n', m)}*{X}_Alim*{a('itc', m)}")), NF_USD, "sum"))
    fcact = lambda m: f"((FC_Desde=0)*({a('pre', m)}+{a('op', m)})+(FC_Desde>=1)*({a('aop', m)}>=FC_Desde)*{a('op', m)})"
    fcusd = lambda m: f"FC_{O}*((FC_Mon=\"PYG\")*{a('itc', m)}+(FC_Mon<>\"PYG\"))*(1+FC_g)^({a('ya', m)}-1)*{fcact(m)}"
    rows += [
        ("pkg_f", "Packaging", "USD", lambda m: f"={X}_Pkg*({a('mp_f', m)}+{a('ins_f', m)}+{a('ov_f', m)}+{a('dir_f', m)})", NF_USD, "sum"),
        ("vs_f", "Seguro de caución, comisiones y logística de venta", "USD", lambda m: f"=({X}_SegV+{X}_ComV)*{a('vtas', m)}", NF_USD, "sum"),
        ("trib_f", "Tributo de maquila", "USD", lambda m: f"={X}_Trib*{a('vtas', m)}", NF_USD, "sum"),
        ("fp_f", "Costos fijos de producción", "USD", lambda m: f'=SUMPRODUCT((FC_Clase="Producción")*{fcusd(m)})', NF_USD, "sum"),
        ("fa_f", "Costos fijos de administración", "USD", lambda m: f'=SUMPRODUCT((FC_Clase="Administración")*{fcusd(m)})', NF_USD, "sum"),
        ("fab_f", "Costos fijos de administración en base de imprevistos", "USD", lambda m: f'=SUMPRODUCT((FC_Clase="Administración")*FC_Base*{fcusd(m)})', NF_USD, "sum"),
        ("impp_f", "Imprevistos de producción", "USD",
         lambda m: (f"={X}_ImpP*({a('mp_f', m)}+{a('ins_f', m)}+{a('ov_f', m)}+{a('desp_f', m)}+{a('pkg_f', m)}+{a('vs_f', m)}"
                    f"+{a('dir_f', m)}+{a('ind_f', m)}+{a('fp_f', m)})"), NF_USD, "sum"),
        ("impa_f", "Imprevistos de administración", "USD", lambda m: f"={X}_ImpA*({a('adm_f', m)}+{a('fab_f', m)})", NF_USD, "sum"),
        ("ivac_f", "IVA contenido en costos fijos", "USD", lambda m: f"=SUMPRODUCT({fcusd(m)}*FC_IVA/(1+FC_IVA))", NF_USD, "sum"),
    ]
    rows.append(("sec", "Estado de resultados"))
    rows += [
        ("MAT", "Materia prima e insumos", "USD", lambda m: R("MAT", m, f"{a('mp_f', m)}+{a('ins_f', m)}"), NF_USD, "sum"),
        ("OVAR", "Otros costos variables", "USD", lambda m: R("OVAR", m, f"{a('ov_f', m)}+{a('desp_f', m)}+{a('pkg_f', m)}+{a('vs_f', m)}"), NF_USD, "sum"),
        ("PERP", "Personal de producción", "USD", lambda m: R("PERP", m, f"{a('dir_f', m)}+{a('ind_f', m)}"), NF_USD, "sum"),
        ("FIJP", "Costos fijos de producción e imprevistos", "USD", lambda m: R("FIJP", m, f"{a('fp_f', m)}+{a('impp_f', m)}"), NF_USD, "sum"),
        ("MB", "MARGEN BRUTO", "USD", lambda m: f"={a('vtas', m)}-{a('MAT', m)}-{a('OVAR', m)}-{a('PERP', m)}-{a('FIJP', m)}", NF_USD, "sum", True),
        ("PERA", "Personal administrativo", "USD", lambda m: R("PERA", m, f"{a('adm_f', m)}"), NF_USD, "sum"),
        ("ADM", "Gastos de administración e imprevistos", "USD", lambda m: R("ADM", m, f"{a('fa_f', m)}+{a('impa_f', m)}"), NF_USD, "sum"),
        ("TRIB", "Tributo de maquila", "USD", lambda m: R("TRIB", m, f"{a('trib_f', m)}"), NF_USD, "sum"),
        ("IVAR", "(+) IVA recuperable (reduce costos)", "USD",
         lambda m: (f"=IF({A('real', m)}=1,0,{a('ivac_f', m)}*{X}_IVArec)" if real else f"={a('ivac_f', m)}*{X}_IVArec"), NF_USD, "sum"),
        ("EBITDA", "EBITDA", "USD", lambda m: f"={a('MB', m)}-{a('PERA', m)}-{a('ADM', m)}-{a('TRIB', m)}+{a('IVAR', m)}", NF_USD, "sum", True),
        ("DEP", "Depreciación y amortización", "USD",
         lambda m: (f"=SUMPRODUCT({'Inv_D' if X == 'P' else '(Inv_CF*' + A('real', m) + '+Inv_BaseV*(1-' + A('real', m) + '))'}"
                    f"*(Inv_F>0)*({m}>=Inv_Svc{X})*({m}<Inv_Svc{X}+Inv_F*12)/(Inv_F*12+(Inv_F=0)))*{hor(m)}"),
         NF_USD, "sum"),
        ("EBIT", "EBIT", "USD", lambda m: f"={a('EBITDA', m)}-{a('DEP', m)}", NF_USD, "sum"),
        ("INT", "Intereses, IVA y comisiones", "USD",
         lambda m: R("INT", m, f"({a('L1_int', m)}+{a('L2_int', m)}+{a('L3_int', m)}+{a('LCI', m)})*(1+{X}_IVAint)"), NF_USD, "sum"),
        ("IRE", "Impuesto a la renta", "USD", lambda m: f"={X}_IRE*({a('EBIT', m)}-{a('INT', m)})", NF_USD, "sum"),
        ("NI", "RESULTADO NETO", "USD", lambda m: f"={a('EBIT', m)}-{a('INT', m)}-{a('IRE', m)}", NF_USD, "sum", True),
    ]

    def bal(code, calc):
        def f(m):
            base = f"IF(AND({A('real', m)}=1,{rgiven(code, m)}),{rv(code, m)},{calc(m)})" if real else calc(m)
            return f"=IF({m}>Fin,{prv(K(code_key[code]), m)},{base})"
        return f
    code_key = {"S_CXC": "cxc", "S_INV": "inv", "S_IVA": "ivab", "S_PROV": "prov"}

    def dias(k, m):
        """Días de cobro/inventario: en meses reales sin saldo cargado no se aplica el escenario."""
        return f"IF({A('real', m)}=1,N(vg_{k}),{X}_{k})" if real else f"{X}_{k}"
    rows.append(("sec", "Capital de trabajo (saldos por días; en el mes de corte se usan los saldos reales)"))
    rows += [
        ("cxc", "Cuentas por cobrar", "USD", bal("S_CXC", lambda m: f"{a('vtas', m)}*{dias('DSO', m)}/DiasMes"), NF_USD, "last"),
        ("inv", "Inventarios (materiales y producto terminado)", "USD",
         bal("S_INV", lambda m: f"{a('MAT', m)}*{dias('DInv', m)}/DiasMes+({a('MAT', m)}+{a('OVAR', m)}+{a('PERP', m)}+{a('FIJP', m)})*{X}_DPT/DiasMes"), NF_USD, "last"),
        ("ivab", "Crédito fiscal IVA", "USD", bal("S_IVA", lambda m: f"{a('IVAR', m)}*{X}_IVAmeses"), NF_USD, "last"),
        ("prov", "Proveedores", "USD",
         bal("S_PROV", lambda m: f"{a('MAT', m)}*{X}_DPOmp/DiasMes+({a('OVAR', m)}+{a('FIJP', m)}+{a('ADM', m)}+{a('TRIB', m)})*{X}_DPOot/DiasMes"),
         NF_USD, "last"),
        ("ctn", "CAPITAL DE TRABAJO NETO", "USD", lambda m: f"={a('cxc', m)}+{a('inv', m)}+{a('ivab', m)}-{a('prov', m)}", NF_USD, "max", True),
        ("dct", "Variación del capital de trabajo", "USD", lambda m: f"={a('ctn', m)}-{prv(K('ctn'), m)}", NF_USD, "sum"),
    ]
    rows.append(("sec", "Flujo de caja e inversión"))
    capex_f = (lambda m: f"(SUMIF(Inv_Q,{m},Inv_R)+SUMIF(Inv_Q,{m},Inv_M)*(1+{X}_dCapex))*{hor(m)}") if real else (lambda m: f"SUMIF(Inv_E,{m},Inv_D)*{hor(m)}")
    rows += [
        ("CFO", "FLUJO OPERATIVO (EBITDA - impuesto - Δ capital de trabajo)", "USD", lambda m: f"=({a('EBITDA', m)}-{a('IRE', m)}-{a('dct', m)})*{hor(m)}", NF_USD, "sum", True),
        ("CAPEX", "CAPEX pagado", "USD", lambda m: R("CAPEX", m, capex_f(m)), NF_USD, "sum"),
        ("AF", "Activo fijo neto (CAPEX acumulado - depreciación acumulada)", "USD",
         lambda m: f"={prv(K('AF'), m)}+{a('CAPEX', m)}-{a('DEP', m)}", NF_USD, "last"),
        ("VT", "Valor terminal (capital de trabajo y activo fijo realizables)", "USD",
         lambda m: f"=IF({m}=Fin,{a('ctn', m)}*{X}_RealCT+{a('AF', m)}*{X}_RealAF,0)", NF_USD, "sum"),
        ("FCFF", "FLUJO LIBRE DEL PROYECTO (FCFF)", "USD", lambda m: f"={a('CFO', m)}-{a('CAPEX', m)}+{a('VT', m)}", NF_USD, "sum", True),
        ("CUML", "Flujo acumulado antes de financiamiento (CFO - CAPEX)", "USD", lambda m: f"={prv(K('CUML'), m)}+{a('CFO', m)}-{a('CAPEX', m)}", NF_USD, "min"),
    ]
    rows.append(("sec", "Préstamos (cronograma contractual)"))
    for j in range(1, NL + 1):
        L = f"{X}_L{j}"
        kk = lambda m, L=L: f"({m}-N({L}_Mes)+1)"
        rows += [
            (f"L{j}_des", f"Préstamo {j}: desembolso", "USD", lambda m, L=L, kk=kk: f"=IF({kk(m)}=1,N({L}_Monto),0)*{hor(m)}", NF_USD, "sum"),
            (f"L{j}_int", f"Préstamo {j}: interés", "USD",
             lambda m, L=L, kk=kk, j=j: f"=IF({kk(m)}>=1,IF({kk(m)}=1,N({L}_Monto),{prv(K(f'L{j}_sal'), m)})*N({L}_Tasa)/12,0)*{hor(m)}", NF_USD, "sum"),
            (f"L{j}_am", f"Préstamo {j}: amortización", "USD",
             lambda m, L=L, kk=kk, j=j: (f"=IF(AND({kk(m)}>N({L}_Gracia),{kk(m)}<=N({L}_Gracia)+N({L}_Cuotas),N({L}_Cuotas)>0),"
                                         f"IF({kk(m)}=N({L}_Gracia)+N({L}_Cuotas),IF({kk(m)}=1,N({L}_Monto),{prv(K(f'L{j}_sal'), m)}),"
                                         f"PMT(N({L}_Tasa)/12,N({L}_Cuotas),-N({L}_Monto))-{a(f'L{j}_int', m)}),0)*{hor(m)}"), NF_USD, "sum"),
            (f"L{j}_sal", f"Préstamo {j}: saldo", "USD",
             lambda m, L=L, kk=kk, j=j: f"=IF({kk(m)}<1,0,IF({kk(m)}=1,N({L}_Monto),{prv(K(f'L{j}_sal'), m)})-{a(f'L{j}_am', m)})", NF_USD, "last"),
        ]
    rows.append(("sec", "Financiamiento y tesorería"))
    LC = f"{X}_LC"
    zr = (lambda m, f: f"=IF({A('real', m)}=1,0,{f})") if real else (lambda m, f: f"={f}")
    rows += [
        ("DESL", "Desembolsos de préstamos", "USD", lambda m: R("DESEMB", m, f"{a('L1_des', m)}+{a('L2_des', m)}+{a('L3_des', m)}"), NF_USD, "sum"),
        ("AML", "Amortización de préstamos (cronograma)", "USD",
         lambda m: R("AMORT", m, f"MAX(0,MIN({prv(K('DEUDA'), m)}-{prv(K('LCS'), m)}+{a('DESL', m)},{a('L1_am', m)}+{a('L2_am', m)}+{a('L3_am', m)}))"),
         NF_USD, "sum"),
        ("APORTE", "Aportes de socios", "USD", lambda m: R("APORTE", m, f"SUMIF(Ap_Mes{O},{m},Ap_Monto{O})*{hor(m)}"), NF_USD, "sum"),
        ("LCA", "Línea rotativa disponible (1/0)", "1/0",
         lambda m: f"=IF(AND(N({LC}_Lim)>0,{m}>=MAX(1,N({LC}_Desde)),{m}<=IF(N({LC}_Hasta)=0,Fin,N({LC}_Hasta))),1,0)", NF_N, None),
        ("LCI", "Línea rotativa: interés del mes (sobre el saldo anterior)", "USD",
         lambda m: zr(m, f"{prv(K('LCS'), m)}*N({LC}_Tasa)/12*{hor(m)}"), NF_USD, "sum"),
        ("CAJA0", "Caja antes de línea rotativa y dividendos", "USD",
         lambda m: (f"={prv(K('CAJA'), m)}+{a('CFO', m)}-{a('CAPEX', m)}+{a('DESL', m)}-{a('AML', m)}-{a('INT', m)}*{hor(m)}+{a('APORTE', m)}"), NF_USD, "min"),
        ("CMIN", "Caja mínima operativa", "USD", lambda m: f"={a('vtas', m)}*{X}_CajaMin/DiasMes*{hor(m)}", NF_USD, None),
        ("LCD", "Línea rotativa: uso", "USD",
         lambda m: zr(m, f"IF({a('LCA', m)}=1,MIN(MAX(0,N({LC}_Lim)-{prv(K('LCS'), m)}),MAX(0,{a('CMIN', m)}-{a('CAJA0', m)})),0)*{hor(m)}"), NF_USD, "sum"),
        ("LCR", "Línea rotativa: devolución (con excedente; total al vencer)", "USD",
         lambda m: zr(m, f"IF({a('LCA', m)}=1,MIN({prv(K('LCS'), m)},MAX(0,{a('CAJA0', m)}-{a('CMIN', m)})),{prv(K('LCS'), m)})*{hor(m)}"), NF_USD, "sum"),
        ("LCS", "Línea rotativa: saldo", "USD",
         lambda m: (f"=IF({A('real', m)}=1,IF({m}=Corte,N(V_LC_Saldo0),0),{prv(K('LCS'), m)}+{a('LCD', m)}-{a('LCR', m)})" if real
                    else f"={prv(K('LCS'), m)}+{a('LCD', m)}-{a('LCR', m)}"), NF_USD, "max"),
        ("DESEMB", "DESEMBOLSOS TOTALES (préstamos + línea)", "USD", lambda m: f"={a('DESL', m)}+{a('LCD', m)}", NF_USD, "sum"),
        ("AMORT", "AMORTIZACIONES TOTALES (préstamos + línea)", "USD", lambda m: f"={a('AML', m)}+{a('LCR', m)}", NF_USD, "sum"),
        ("DEUDA", "DEUDA FINANCIERA (saldo)", "USD", lambda m: f"={prv(K('DEUDA'), m)}+{a('DESEMB', m)}-{a('AMORT', m)}", NF_USD, "max", True),
        ("CAJA1", "Caja después de la línea rotativa", "USD", lambda m: f"={a('CAJA0', m)}+{a('LCD', m)}-{a('LCR', m)}", NF_USD, "min"),
        ("NIY", "Resultado neto acumulado del año", "USD",
         lambda m: f"=IF({A('mes', m)}=1,{a('NI', m)},{prv(K('NIY'), m)}+{a('NI', m)})", NF_USD, None),
        ("NIP", "Resultado neto del año anterior", "USD",
         lambda m: "=0" if m == 1 else f"=IF({A('mes', m)}=1,{prv(K('NIY'), m)},{prv(K('NIP'), m)})", NF_USD, None),
        ("DIVD", "Dividendos según política (resultado del año anterior)", "USD",
         lambda m: (f"=IF(AND({A('mes', m)}=MAX(1,{X}_MesDiv),{A('anio', m)}>=2),INDEX(Payout_{O},1,MIN({NY},{A('anio', m)}-1))"
                    f"*MAX(0,{a('NIP', m)}),0)*{hor(m)}"), NF_USD, "sum"),
        ("DIV", "Dividendos pagados (solo con caja sobre el mínimo)", "USD",
         lambda m: R("DIV", m, f"MIN({a('DIVD', m)},MAX(0,{a('CAJA1', m)}-{a('CMIN', m)}))"), NF_USD, "sum"),
        ("CAJA", "CAJA FINAL", "USD",
         lambda m: (f"=IF(AND({A('real', m)}=1,{rgiven('S_CAJA', m)}),{rv('S_CAJA', m)},{a('CAJA1', m)}-{a('DIV', m)})" if real
                    else f"={a('CAJA1', m)}-{a('DIV', m)}"), NF_USD, "min", True),
        ("AJ", "Ajuste de conciliación con la caja real", "USD", lambda m: f"={a('CAJA', m)}-({a('CAJA1', m)}-{a('DIV', m)})", NF_USD, "sum"),
        ("BRECHA", "Brecha: caja por debajo de la caja mínima", "USD", lambda m: f"=MAX(0,{a('CMIN', m)}-{a('CAJA', m)})*{hor(m)}", NF_USD, "max", True),
        ("DEF", "Caja negativa a cubrir por los socios (aporte implícito acumulado)", "USD", lambda m: f"=MAX(0,-{a('CAJA', m)})", NF_USD, "max"),
    ]
    rows.append(("sec", "Valuación y control"))
    rows += [
        ("DFP", "Factor de descuento del proyecto", "", lambda m: f"=1/(1+{X}_WACC)^(({m}-1)/12)", NF_USD4, None),
        ("DFE", "Factor de descuento de los socios", "", lambda m: f"=1/(1+{X}_Ke)^(({m}-1)/12)", NF_USD4, None),
        ("EQ", "Flujo de los socios (-aportes - cobertura de caja negativa + dividendos + valor final)", "USD",
         lambda m: (f"=-{a('APORTE', m)}-({a('DEF', m)}-{prv(K('DEF'), m)})+{a('DIV', m)}"
                    f"+IF({m}=Fin,{a('CAJA', m)}+{a('DEF', m)}+{a('VT', m)}-{a('DEUDA', m)},0)"), NF_USD, "sum"),
        ("CUMF", "FCFF acumulado", "USD", lambda m: f"={prv(K('CUMF'), m)}+{a('FCFF', m)}", NF_USD, "last"),
        ("PB", "Recupero de la inversión (mes con interpolación)", "mes",
         lambda m: "=0" if m == 1 else f"=IF(AND({prv(K('CUMF'), m)}<0,{a('CUMF', m)}>=0),{m - 2}+(-{prv(K('CUMF'), m)})/{a('FCFF', m)},0)", NF_USD2, "max"),
        ("PAT", "Patrimonio (aportes + resultados - dividendos + ajustes)", "USD",
         lambda m: f"={prv(K('PAT'), m)}+{a('APORTE', m)}+{a('NI', m)}-{a('DIV', m)}+{a('AJ', m)}", NF_USD, "last"),
        ("CHK", "CONTROL: activo - pasivo - patrimonio (0 = OK)", "USD",
         lambda m: (f"=ROUND({a('CAJA', m)}+{a('cxc', m)}+{a('inv', m)}+{a('ivab', m)}+{a('AF', m)}-{a('prov', m)}-{a('DEUDA', m)}-{a('PAT', m)},2)"),
         NF_USD2, "maxabs", True),
    ]
    return rows


def build_calculo():
    ws = WS["Calculo"]
    title(ws, "Cálculo — motor mensual (no editar)",
          "Bloque P = presupuesto original (solo supuestos originales). Bloque V = vigente: Reales hasta el último mes cerrado + proyección con escenario. Mismas fórmulas.")
    widths(ws, {"A": 9, "B": 60, "C": 7, "D": 3, "E": 15})
    for m in range(1, N + 1):
        ws.column_dimensions[mc(m)].width = 11
    # asignación de filas
    r = 4
    cal = [("m", "Mes n°", lambda m: m, NF_N), ("fecha", "Fecha", lambda m: f'=IF(n_Fecha1="","",EOMONTH(n_Fecha1,{m}-1))', NF_MES),
           ("anio", "Año del modelo", lambda m: f"=ROUNDUP({m}/12,0)", NF_N), ("mes", "Mes del año del modelo", lambda m: f"=MOD({m}-1,12)+1", NF_N),
           ("real", "Mes real cerrado (1/0)", lambda m: f"=IF({m}<=Corte,1,0)", NF_N), ("hor", "Dentro del horizonte (1/0)", lambda m: f"=IF({m}<=Fin,1,0)", NF_N),
           ("tcr", "TC para convertir reales (real del mes o vigente)", None, NF_INT)]
    for k, *_ in cal:
        ROW[k] = r
        r += 1
    r += 1
    blocks = {}
    for X in ("P", "V"):
        rows = version_rows(X)
        blocks[X] = (r, rows)
        r += 1   # título del bloque
        for it in rows:
            if it[0] == "sec":
                r += 1
            else:
                ROW[f"{X}.{it[0]}"] = r
                r += 1
        r += 2
    # escritura del calendario
    for k, lab, fn, fmt in cal:
        rr = ROW[k]
        put(ws, rr, 2, lab, font=F_H if k == "m" else F_B, fill=FL_H if k == "m" else None)
        put(ws, rr, 1, k, font=F_S)
        for m in range(1, N + 1):
            if k == "tcr":
                v = f"=IF(N(Reales!{mc(m)}{RR['TC']})>0,Reales!{mc(m)}{RR['TC']},INDEX(TC_V,1,MIN({NY},{A('anio', m)})))"
            else:
                v = fn(m)
            put(ws, rr, FCOL + m - 1, v, fmt=fmt, font=F_H if k == "m" else F, fill=FL_H if k == "m" else None)
    # bloques
    for X in ("P", "V"):
        start, rows = blocks[X]
        lab = ("BLOQUE P — PRESUPUESTO ORIGINAL" if X == "P" else "BLOQUE V — VIGENTE: REAL HASTA EL CORTE + PROYECCIÓN")
        for c in range(1, 6):
            ws.cell(row=start, column=c).fill = FL_P if X == "P" else FL_V
        put(ws, start, 2, lab, font=F_SEC, fill=FL_P if X == "P" else FL_V)
        r = start + 1
        for it in rows:
            if it[0] == "sec":
                put(ws, r, 2, it[1], font=F_SEC, fill=FL_SEC)
                for c in range(1, 6):
                    ws.cell(row=r, column=c).fill = FL_SEC
                r += 1
                continue
            key, lab, unit, fn, fmt, tot = it[:6]
            bold = len(it) > 6 and it[6]
            rr = ROW[f"{X}.{key}"]
            assert rr == r, (X, key, rr, r)
            put(ws, r, 1, f"{X}.{key}", font=F_S)
            put(ws, r, 2, lab, font=F_B if bold else (F_LINK if str(lab).startswith("=") else F), fill=FL_TOT if bold else None)
            put(ws, r, 3, unit)
            rng = f"F{r}:{MN}{r}"
            totf = {"sum": f"=SUM({rng})", "last": f"=INDEX({rng},1,MIN({N},Fin))", "max": f"=MAX({rng})", "min": f"=MIN({rng})",
                    "maxabs": f"=MAX(ABS(MAX({rng})),ABS(MIN({rng})))"}.get(tot)
            if totf:
                put(ws, r, 5, totf, fmt=fmt, font=F_B, fill=FL_TOT)
            for m in range(1, N + 1):
                put(ws, r, FCOL + m - 1, fn(m), fmt=fmt, font=F_B if bold else F)
            r += 1
    ws.freeze_panes = f"F{ROW['real'] + 1}"


# ============================================================================ RESULTADOS
RS = {}


def build_resultados():
    ws = WS["Resultados"]
    CUR[0] = "Resultados"
    title(ws, "Resultados — presupuesto original vs vigente (real + proyección)")
    put(ws, 2, 1, '="Escenario de proyección: "&n_Escenario&"   |   Último mes cerrado: "&IF(Corte=0,"sin datos reales","mes "&Corte)', font=F_AL)
    widths(ws, {"A": 6, "B": 56, "C": 16, "D": 16, "E": 16, "F": 60})
    for c in range(4, 16):
        ws.column_dimensions[CL(c)].width = 13
    YR = f"Calculo!{RG('anio')}"
    r = 4
    section(ws, r, "A. Indicadores clave", 15); r += 1
    header(ws, r, ["", "Indicador", "Presupuesto original", "Vigente (real + proyección)", "Desvío", "Lectura"], 1, 30)
    r += 1

    def g(X, k):
        return f"Calculo!{RG(X + '.' + k)}"

    def tot(X, k):
        return f"Calculo!$E${ROW[X + '.' + k]}"

    def irr(X, k):
        rg = g(X, k)
        return (f'IF(COUNTIF({rg},"<0")=0,"No existe",IF(COUNTIF({rg},">0")=0,"No existe",'
                f'IFERROR((1+IRR({rg},0.01))^12-1,IFERROR((1+IRR({rg},-0.01))^12-1,"No converge"))))')

    kpis = [
        ("v1", "Ventas del año 1 del modelo", lambda X: f"=SUMIFS({g(X, 'vtas')},{YR},1)", NF_USD, ""),
        ("vt", "Ventas totales del horizonte", lambda X: f"={tot(X, 'vtas')}", NF_USD, ""),
        ("e1", "EBITDA del año 1 del modelo", lambda X: f"=SUMIFS({g(X, 'EBITDA')},{YR},1)", NF_USD, ""),
        ("et", "EBITDA total del horizonte", lambda X: f"={tot(X, 'EBITDA')}", NF_USD, ""),
        ("em", "Margen EBITDA del horizonte", lambda X: f"=IF({tot(X, 'vtas')}=0,\"\",{tot(X, 'EBITDA')}/{tot(X, 'vtas')})", NF_PCT, ""),
        ("nt", "Resultado neto total", lambda X: f"={tot(X, 'NI')}", NF_USD, ""),
        ("capex", "Inversión total (CAPEX)", lambda X: f"={tot(X, 'CAPEX')}", NF_USD, "Vigente: real pagado + pendiente con escenario."),
        ("nec", "Necesidad máxima de fondos antes de financiamiento", lambda X: f"=MAX(0,-MIN({g(X, 'CUML')}))", NF_USD,
         "Inversión + capital de trabajo + pérdidas iniciales que deben financiar socios y bancos."),
        ("necm", "Mes de la necesidad máxima", lambda X: f'=IF(MIN({g(X, "CUML")})>=0,"",MATCH(MIN({g(X, "CUML")}),{g(X, "CUML")},0))', NF_N, ""),
        ("ap", "Aportes de socios", lambda X: f"={tot(X, 'APORTE')}", NF_USD, ""),
        ("apimp", "Aporte adicional necesario para no tener caja negativa", lambda X: f"={tot(X, 'DEF')}", NF_USD,
         "Fondos no previstos que deberían aportar socios o bancos (máximo saldo de caja negativa)."),
        ("deuda", "Deuda financiera máxima", lambda X: f"={tot(X, 'DEUDA')}", NF_USD, "Préstamos + línea rotativa."),
        ("lcmax", "Línea rotativa: uso máximo", lambda X: f"={tot(X, 'LCS')}", NF_USD, "0 = no se usó o no hay línea."),
        ("intt", "Intereses, IVA y comisiones totales", lambda X: f"={tot(X, 'INT')}", NF_USD, ""),
        ("canc", "Cancelación total de la deuda (años desde el mes 1)",
         lambda X: (f'=IF(MAX({g(X, "DEUDA")})<1,"Sin deuda",IF(INDEX({g(X, "DEUDA")},1,MIN({N},Fin))>1,"No cancela en el horizonte",'
                    f'(_xlfn.MAXIFS(Calculo!{RG("m")},{g(X, "DEUDA")},">1")+1)/12))'), NF_USD2, "Último mes con saldo de deuda + 1, en años."),
        ("cajamin", "Caja final mínima del horizonte", lambda X: f"=MIN({g(X, 'CAJA')})", NF_USD, "Negativa = financiamiento no previsto."),
        ("brecha", "Brecha máxima bajo la caja mínima", lambda X: f"={tot(X, 'BRECHA')}", NF_USD, "No se insertan fondos ficticios: la brecha queda visible."),
        ("brechan", "Meses con brecha", lambda X: f'=SUMPRODUCT(--({g(X, "BRECHA")}>0.5))', NF_N, ""),
        ("dscr", "DSCR mínimo desde el 2.º año de operación (flujo operativo / servicio de deuda)", None, NF_X,
         "Servicio = intereses + cuotas de préstamos (sin devoluciones de la línea). < 1 = la operación no alcanza a pagar la deuda ese año."),
        ("van", "VAN del proyecto (FCFF a WACC)", lambda X: f"=SUMPRODUCT({g(X, 'FCFF')},{g(X, 'DFP')})", NF_USD, "Excluye préstamos, intereses, aportes y dividendos."),
        ("tir", "TIR del proyecto (anual)", lambda X: "=" + irr(X, "FCFF"), NF_PCT2, ""),
        ("vane", "VAN de los socios (a Ke)", lambda X: f"=SUMPRODUCT({g(X, 'EQ')},{g(X, 'DFE')})", NF_USD,
         "Aportes, dividendos pagados y valor final. La caja negativa se considera cubierta por los socios hasta que se recupera."),
        ("tire", "TIR de los socios (anual)", lambda X: "=" + irr(X, "EQ"), NF_PCT2, ""),
        ("pb", "Recupero de la inversión del proyecto (años)", lambda X: f'=IF(MAX({g(X, "PB")})=0,"No recupera",MAX({g(X, "PB")})/12)', NF_USD2, ""),
        ("chk", "Balance de control (0 = OK)", lambda X: f"={tot(X, 'CHK')}", NF_USD2, ""),
    ]
    for key, lab, fn, fmt, note in kpis:
        RS["k_" + key] = r
        put(ws, r, 2, lab)
        for col, X in [(3, "P"), (4, "V")]:
            if fn:
                put(ws, r, col, fn(X), fmt=fmt, fill=FL_P if X == "P" else FL_V, font=F_B)
        put(ws, r, 5, f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),D{r}-C{r},"")', fmt=fmt)
        put(ws, r, 6, note, font=F_S)
        r += 1
    nm("r_BrechaV", "Resultados", f"D{RS['k_brecha']}")
    nm("r_ChkP", "Resultados", f"C{RS['k_chk']}")
    nm("r_ChkV", "Resultados", f"D{RS['k_chk']}")
    r += 1
    put(ws, r, 2, "Plan publicado (referencia): VAN / TIR", font=F_S)
    put(ws, r, 3, '=IF(pub_VAN="","",pub_VAN)', fmt=NF_USD)
    put(ws, r, 4, '=IF(pub_TIR="","",pub_TIR)', fmt=NF_PCT2)
    put(ws, r, 6, "Cifras del documento original, calculadas con su propia metodología.", font=F_S)
    r += 2

    # ---------------------------------------------------------- tablas anuales
    def yhdr(rr, lab):
        header(ws, rr, ["", lab, "Total"] + [f"Año {y}" for y in range(1, NYR + 1)], 1)
        put(ws, rr + 1, 2, "Estado del año", font=F_S)
        put(ws, rr + 2, 2, "Año calendario", font=F_S)
        for y in range(1, NYR + 1):
            c = 3 + y
            put(ws, rr + 1, c, f'=IF(Corte>={12 * y},"Real",IF(Corte>{12 * (y - 1)},"Real+Proy.","Proyección"))', font=F_S, al=CEN)
            put(ws, rr + 2, c, f'=IF(n_Fecha1="","",YEAR(EOMONTH(n_Fecha1,{12 * (y - 1)})))', font=F_S, fmt=NF_N, al=CEN)
        return rr + 3

    def arow(rr, key, lab, X, src, kind="sum", fmt=NF_USD, bold=False):
        RS[key] = rr
        put(ws, rr, 1, key, font=F_S)
        put(ws, rr, 2, lab, font=F_B if bold else F, fill=FL_TOT if bold else None)
        for y in range(1, NYR + 1):
            c = 3 + y
            if kind == "sum":
                f = f"=SUMIFS({g(X, src)},{YR},{y})"
            elif kind == "end":
                f = f"=INDEX({g(X, src)},1,MIN({N},{12 * y}))"
            elif kind == "max":
                f = f"=_xlfn.MAXIFS({g(X, src)},{YR},{y})"
            else:
                f = src(y, CL(c))
            put(ws, rr, c, f, fmt=fmt, font=F_B if bold else F)
        if kind == "sum":
            put(ws, rr, 3, f"=SUM(D{rr}:O{rr})", fmt=fmt, font=F_B, fill=FL_TOT)
        elif kind in ("end",):
            put(ws, rr, 3, f"=INDEX({g(X, src)},1,MIN({N},Fin))", fmt=fmt, font=F_B, fill=FL_TOT)
        elif kind == "max":
            put(ws, rr, 3, f"=MAX(D{rr}:O{rr})", fmt=fmt, font=F_B, fill=FL_TOT)
        return rr + 1

    section(ws, r, "B. Vigente (real hasta el corte + proyección) — resumen anual", 15); r += 1
    r = yhdr(r, "Concepto (USD)")
    lines = [("V_vtas", "Ventas", "vtas", "sum", True), ("V_MAT", "Materia prima e insumos", "MAT", "sum", False),
             ("V_OVAR", "Otros costos variables", "OVAR", "sum", False), ("V_PERP", "Personal de producción", "PERP", "sum", False),
             ("V_FIJP", "Costos fijos de producción e imprevistos", "FIJP", "sum", False), ("V_MB", "Margen bruto", "MB", "sum", True),
             ("V_PERA", "Personal administrativo", "PERA", "sum", False), ("V_ADM", "Gastos de administración", "ADM", "sum", False),
             ("V_TRIB", "Tributo de maquila", "TRIB", "sum", False), ("V_IVAR", "(+) IVA recuperable", "IVAR", "sum", False),
             ("V_EBITDA", "EBITDA", "EBITDA", "sum", True), ("V_DEP", "Depreciación", "DEP", "sum", False), ("V_INT", "Intereses, IVA y comisiones", "INT", "sum", False),
             ("V_IRE", "Impuesto a la renta", "IRE", "sum", False), ("V_NI", "Resultado neto", "NI", "sum", True),
             ("V_CFO", "Flujo operativo", "CFO", "sum", True), ("V_CAPEX", "CAPEX pagado", "CAPEX", "sum", False), ("V_FCFF", "FCFF (incluye valor terminal)", "FCFF", "sum", True),
             ("V_APORTE", "Aportes de socios", "APORTE", "sum", False), ("V_DESEMB", "Desembolsos de préstamos", "DESEMB", "sum", False),
             ("V_AMORT", "Amortizaciones (préstamos + línea)", "AMORT", "sum", False), ("V_AML", "   de las cuales: cuotas de préstamos", "AML", "sum", False),
             ("V_DIV", "Dividendos pagados", "DIV", "sum", False),
             ("V_CAJA", "Caja al cierre", "CAJA", "end", True), ("V_DEUDA", "Deuda al cierre", "DEUDA", "end", False), ("V_LCS", "   de la cual: línea rotativa", "LCS", "end", False),
             ("V_CTN", "Capital de trabajo neto al cierre", "ctn", "end", False), ("V_BRECHA", "Brecha máxima del año", "BRECHA", "max", False)]
    for key, lab, src, kind, bold in lines:
        r = arow(r, key, lab, "V", src, kind, bold=bold)
    r = arow(r, "V_mg", "Margen EBITDA", "V", lambda y, c: f'=IF({c}{RS["V_vtas"]}=0,"",{c}{RS["V_EBITDA"]}/{c}{RS["V_vtas"]})', "f", NF_PCT)
    r = arow(r, "V_DS", "Servicio de deuda (intereses + amortización)", "V", lambda y, c: f"={c}{RS['V_INT']}+{c}{RS['V_AML']}", "f")
    r = arow(r, "V_DSCR", "DSCR (flujo operativo / servicio de deuda)", "V",
             lambda y, c: f'=IF({c}{RS["V_DS"]}>1,{c}{RS["V_CFO"]}/{c}{RS["V_DS"]},"")', "f", NF_X)
    r = arow(r, "YN", "Año n° (auxiliar)", "V", lambda y, c: f"={y}", "f", NF_N)
    put(ws, r, 2, "Unidades vendidas por producto", font=F_B)
    r += 1
    for i in range(NP):
        r = arow(r, f"V_u{i + 1}", f'=IF(INDEX(Prod_Nombre,{i + 1})="","Producto {i + 1}",INDEX(Prod_Nombre,{i + 1}))', "V", f"u{i + 1}", "sum", NF_INT)
    r += 1
    section(ws, r, "C. Presupuesto original — resumen anual", 15); r += 1
    r = yhdr(r, "Concepto (USD)")
    for key, lab, src, kind, bold in [("P_vtas", "Ventas", "vtas", "sum", True), ("P_EBITDA", "EBITDA", "EBITDA", "sum", True), ("P_NI", "Resultado neto", "NI", "sum", True),
                                      ("P_CFO", "Flujo operativo", "CFO", "sum", False), ("P_CAPEX", "CAPEX", "CAPEX", "sum", False),
                                      ("P_FCFF", "FCFF (incluye valor terminal)", "FCFF", "sum", True), ("P_APORTE", "Aportes de socios", "APORTE", "sum", False),
                                      ("P_INT", "Intereses, IVA y comisiones", "INT", "sum", False), ("P_AMORT", "Amortizaciones (préstamos + línea)", "AMORT", "sum", False),
                                      ("P_AML", "   de las cuales: cuotas de préstamos", "AML", "sum", False),
                                      ("P_CAJA", "Caja al cierre", "CAJA", "end", True), ("P_DEUDA", "Deuda al cierre", "DEUDA", "end", False),
                                      ("P_CTN", "Capital de trabajo neto al cierre", "ctn", "end", False)]:
        r = arow(r, key, lab, "P", src, kind, bold=bold)
    r = arow(r, "P_DS", "Servicio de deuda", "P", lambda y, c: f"={c}{RS['P_INT']}+{c}{RS['P_AML']}", "f")
    r = arow(r, "P_DSCR", "DSCR", "P", lambda y, c: f'=IF({c}{RS["P_DS"]}>1,{c}{RS["P_CFO"]}/{c}{RS["P_DS"]},"")', "f", NF_X)
    put(ws, r, 2, "Unidades vendidas por producto", font=F_B)
    r += 1
    for i in range(NP):
        r = arow(r, f"P_u{i + 1}", f'=IF(INDEX(Prod_Nombre,{i + 1})="","Producto {i + 1}",INDEX(Prod_Nombre,{i + 1}))', "P", f"u{i + 1}", "sum", NF_INT)
    r += 1
    section(ws, r, "D. Desvíos: vigente - presupuesto original", 15); r += 1
    r = yhdr(r, "Concepto (USD)")
    for key, lab, kv, kp in [("D_vtas", "Ventas", "V_vtas", "P_vtas"), ("D_EBITDA", "EBITDA", "V_EBITDA", "P_EBITDA"), ("D_NI", "Resultado neto", "V_NI", "P_NI"),
                             ("D_CAPEX", "CAPEX", "V_CAPEX", "P_CAPEX"), ("D_FCFF", "FCFF", "V_FCFF", "P_FCFF"), ("D_CAJA", "Caja al cierre", "V_CAJA", "P_CAJA")]:
        RS[key] = r
        put(ws, r, 2, lab)
        for y in range(0, NYR + 1):
            c = CL(3 + y)
            put(ws, r, 3 + y, f"={c}{RS[kv]}-{c}{RS[kp]}", fmt=NF_USD)
        r += 1
    RS["D_u"] = r
    put(ws, r, 2, "Unidades totales")
    for y in range(0, NYR + 1):
        c = CL(3 + y)
        put(ws, r, 3 + y, f"=SUM({c}{RS['V_u1']}:{c}{RS['V_u' + str(NP)]})-SUM({c}{RS['P_u1']}:{c}{RS['P_u' + str(NP)]})", fmt=NF_INT)
    r += 1
    put(ws, r, 2, "% de desvío en ventas")
    for y in range(0, NYR + 1):
        c = CL(3 + y)
        put(ws, r, 3 + y, f'=IF({c}{RS["P_vtas"]}=0,"",{c}{RS["D_vtas"]}/{c}{RS["P_vtas"]})', fmt=NF_PCT)
    r += 2
    section(ws, r, "E. Plan publicado (referencia) vs presupuesto original del modelo — explica diferencias de metodología", 15); r += 1
    r = yhdr(r, "Concepto (USD)")
    for k, lab, kp in [("Ventas", "Ventas publicadas", "P_vtas"), ("EBITDA", "EBITDA publicado", "P_EBITDA"), ("Utilidad", "Resultado neto publicado", "P_NI")]:
        RS["pub_" + k] = r
        put(ws, r, 2, lab)
        for y in range(1, NY + 1):
            put(ws, r, 3 + y, f'=IF(INDEX(pub_{k},1,{y})="","",INDEX(pub_{k},1,{y}))', fmt=NF_USD, font=F_LINK)
        put(ws, r, 3, f"=SUM(D{r}:M{r})", fmt=NF_USD, font=F_B, fill=FL_TOT)
        r += 1
        put(ws, r, 2, f"   Diferencia presupuesto del modelo - publicado")
        for y in range(1, NY + 1):
            c = CL(3 + y)
            put(ws, r, 3 + y, f'=IF({c}{r - 1}="","",{c}{RS[kp]}-{c}{r - 1})', fmt=NF_USD)
        r += 1
    # DSCR mínimo
    for col, k in [(3, "P_DSCR"), (4, "V_DSCR")]:
        crit = f'D{RS[k[0] + "_DS"]}:O{RS[k[0] + "_DS"]},">1",D{RS["YN"]}:O{RS["YN"]},">="&(ROUNDUP({k[0]}_Inicio/12,0)+1)'
        put(ws, RS["k_dscr"], col, f'=IF(COUNTIFS({crit})=0,"Sin deuda",_xlfn.MINIFS(D{RS[k]}:O{RS[k]},{crit}))',
            fmt=NF_X, fill=FL_P if col == 3 else FL_V, font=F_B)
    add_charts(ws, r + 2)
    ws.freeze_panes = "C6"


def add_charts(ws, anchor):
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.series import SeriesLabel
    lab_row = anchor
    for y in range(1, NYR + 1):
        ws.cell(row=lab_row, column=3 + y, value=f"Año {y}")
    put(ws, lab_row, 2, "Etiquetas de gráficos", font=F_S)
    cats = Reference(ws, min_col=4, max_col=3 + NYR, min_row=lab_row)

    def ref(key):
        return Reference(ws, min_col=4, max_col=3 + NYR, min_row=RS[key])

    ch = BarChart()
    ch.title = "Ventas: presupuesto vs vigente"
    ch.height, ch.width = 7.5, 17
    for k, n in [("P_vtas", "Presupuesto"), ("V_vtas", "Vigente")]:
        ch.add_data(ref(k), from_rows=True, titles_from_data=False)
        ch.series[-1].tx = SeriesLabel(v=n)
    ch.set_categories(cats)
    ws.add_chart(ch, f"B{anchor + 2}")
    ch2 = LineChart()
    ch2.title = "EBITDA y caja al cierre"
    ch2.height, ch2.width = 7.5, 17
    for k, n in [("P_EBITDA", "EBITDA presupuesto"), ("V_EBITDA", "EBITDA vigente"), ("P_CAJA", "Caja presupuesto"), ("V_CAJA", "Caja vigente")]:
        ch2.add_data(ref(k), from_rows=True, titles_from_data=False)
        ch2.series[-1].tx = SeriesLabel(v=n)
        ch2.series[-1].smooth = False
    ch2.set_categories(cats)
    ch2.x_axis.tickLblPos = "low"
    ws.add_chart(ch2, f"F{anchor + 2}")


# ============================================================================ CONTROL
def build_control():
    ws = WS["Control"]
    title(ws, "Control — validaciones y datos faltantes", "Estado OK cuando el valor está dentro de lo esperado. Revisar cada cierre.")
    widths(ws, {"A": 6, "B": 70, "C": 16, "D": 14, "E": 12, "F": 80})
    r = 4
    header(ws, r, ["Cód.", "Control", "Valor", "Esperado", "Estado", "Qué hacer si no está OK"], 1)
    r += 1
    first = r
    real_m = f"(Calculo!{RG('real')}=1)"
    checks = [
        ("Balance del presupuesto cuadra (activo - pasivo - patrimonio)", "=r_ChkP", "= 0", "abs0", "Error de fórmula: no editar Calculo; revisar entradas en blanco."),
        ("Balance vigente cuadra", "=r_ChkV", "= 0", "abs0", "Idem."),
        ("Meses cerrados sin ventas cargadas", f"=SUMPRODUCT({real_m}*(Calculo!{RG('V.vtas')}=0))", "= 0", "zero",
         "Cargar ventas reales del mes en Reales (si realmente no hubo ventas, cargar 0 explícito y documentarlo)."),
        ("Meses cerrados sin costos de materiales ni personal cargados",
         f"=SUMPRODUCT({real_m}*(Calculo!{RG('V.MAT')}+Calculo!{RG('V.PERP')}+Calculo!{RG('V.PERA')}=0))", "= 0", "zero", "Cargar costos reales del mes."),
        ("Saldo de caja real cargado en el mes de corte", f'=IF(Corte=0,1,IF(ISNUMBER(INDEX(Reales!$F${RR["S_CAJA"]}:${MN}${RR["S_CAJA"]},1,Corte)),1,0))', "= 1", "one",
         "Cargar el saldo de caja y bancos al cierre del último mes cerrado."),
        ("Saldos de clientes, inventarios y proveedores cargados en el mes de corte",
         f'=IF(Corte=0,1,IF(AND(ISNUMBER(INDEX(Reales!$F${RR["S_CXC"]}:${MN}${RR["S_CXC"]},1,Corte)),ISNUMBER(INDEX(Reales!$F${RR["S_INV"]}:${MN}${RR["S_INV"]},1,Corte)),'
         f'ISNUMBER(INDEX(Reales!$F${RR["S_PROV"]}:${MN}${RR["S_PROV"]},1,Corte))),1,0))', "= 1", "one", "Sin estos saldos el capital de trabajo se estima por días."),
        ("Diferencia de deuda: real al corte - calculada",
         f'=IF(Corte=0,0,IF(ISNUMBER(INDEX(Reales!$F${RR["S_DEUDA"]}:${MN}${RR["S_DEUDA"]},1,Corte)),INDEX(Reales!$F${RR["S_DEUDA"]}:${MN}${RR["S_DEUDA"]},1,Corte)'
         f'-INDEX(Calculo!{RG("V.DEUDA")},1,Corte),0))', "= 0", "abs1", "Cargar desembolsos y amortizaciones reales o ajustar las condiciones vigentes del préstamo."),
        ("CAPEX: pagado por ítem (Inversion) - CAPEX real mensual (Reales)", f"=i_Pagado-SUM(Reales!$F${RR['CAPEX']}:${MN}${RR['CAPEX']})", "= 0", "abs1",
         "Ambos registros deben conciliar."),
        ("Tipo de cambio cargado en todos los años (presupuesto y vigente)", "=COUNTIF(TC_O,\">0\")+COUNTIF(TC_V,\">0\")", f"= {2 * NY}", "tc",
         "Completar el TC por año en Supuestos B (los costos en PYG se convierten con él)."),
        ("Productos con volumen vigente por encima de la capacidad (años)", "=SUM(c_CapExceso)", "= 0", "zero",
         "Revisar volumen o registrar la ampliación de capacidad (Productos A)."),
        ("Productos con volumen y sin precio vigente", "=SUMPRODUCT((Precio_V=0)*(VolTot_V>0))", "= 0", "zero", "Completar el precio."),
        ("Préstamos del presupuesto cancelados dentro del horizonte", "=r_DeudaFinP", "= 0", "abs1", "Revisar plazo o el horizonte."),
        ("Brecha de caja vigente (meses bajo la caja mínima)", f'=SUMPRODUCT(--(Calculo!{RG("V.BRECHA")}>0.5))', "= 0", "zero",
         "Definir aportes, financiamiento o mejoras de capital de trabajo (alerta, no error de fórmula)."),
        ("Negocio en operación: mes real de inicio y último mes cerrado cargados", '=IF(n_Etapa="Operación",IF(AND(ISNUMBER(n_InicioReal),Corte>0),1,0),1)',
         "= 1", "one", "Completar en Inicio el mes real de inicio de operación y cargar los meses cerrados en Reales."),
        ("Horizonte dentro de la rejilla de 144 meses (mes de inicio ≤ 25 con 10 años)", f"=IF(MAX(Fin,V_Inicio)<={N},1,0)", "= 1", "one",
         "Reducir los años evaluados o el mes de inicio (la rejilla tiene 144 meses)."),
        ("Supuestos vigentes modificados respecto del presupuesto (informativo)", "=c_Modificados", "≥ 0", "info", "Documentar en la nota de cada supuesto el motivo del cambio."),
    ]
    for i, (lab, f, exp, kind, todo) in enumerate(checks):
        put(ws, r, 1, f"C{i + 1:02d}")
        put(ws, r, 2, lab)
        put(ws, r, 3, f, fmt=NF_USD2)
        put(ws, r, 4, exp, al=CEN)
        cond = {"abs0": f"ABS(C{r})<0.01", "abs1": f"ABS(C{r})<1", "zero": f"C{r}=0", "one": f"C{r}=1", "tc": f"C{r}={2 * NY}", "info": "TRUE"}[kind]
        put(ws, r, 5, f'=IF({cond},"OK","REVISAR")', font=F_B, al=CEN)
        put(ws, r, 6, todo, font=F_S)
        r += 1
    last = r - 1
    ws.conditional_formatting.add(f"E{first}:E{last}", CellIsRule(operator="equal", formula=['"OK"'], fill=FL_OK))
    ws.conditional_formatting.add(f"E{first}:E{last}", CellIsRule(operator="equal", formula=['"REVISAR"'], fill=FL_BAD))
    r += 1
    put(ws, r, 2, "Controles en REVISAR", font=F_B)
    put(ws, r, 3, f'=COUNTIF(E{first}:E{last},"REVISAR")', fmt=NF_N, font=F_B)
    nm("c_Alertas", "Control", f"C{r}")
    r += 1
    put(ws, r, 2, "Celdas vigentes sobrescritas (difieren del presupuesto)")
    put(ws, r, 3, (f"=SUMPRODUCT(--(Supuestos!$E$6:$E${SUPR['p_last']}<>Supuestos!$D$6:$D${SUPR['p_last']}))"
                   f"+SUMPRODUCT(--(Precio_V<>Precio_O))+SUMPRODUCT(--(MP_V<>MP_O))+SUMPRODUCT(--(INS_V<>INS_O))+SUMPRODUCT(--(OV_V<>OV_O))"
                   f"+SUMPRODUCT(--(Vol_V<>Vol_O))+SUMPRODUCT(--(Sal_V<>Sal_O))+SUMPRODUCT(--(Dot_V<>Dot_O))+SUMPRODUCT(--(FC_V<>FC_O))"
                   f"+SUMPRODUCT(--(TC_V<>TC_O))+SUMPRODUCT(--(dPrecio_V<>dPrecio_O))+SUMPRODUCT(--(dMP_V<>dMP_O))+SUMPRODUCT(--(dOV_V<>dOV_O))"
                   f"+SUMPRODUCT(--(dSal_V<>dSal_O))+SUMPRODUCT(--(Payout_V<>Payout_O))+SUMPRODUCT(--(Ap_MesV<>Ap_MesO))+SUMPRODUCT(--(Ap_MontoV<>Ap_MontoO))"),
        fmt=NF_N, font=F_B)
    nm("c_Modificados", "Control", f"C{r}")
    r += 1
    put(ws, r, 2, "Deuda del presupuesto al final del horizonte")
    put(ws, r, 3, f"=INDEX(Calculo!{RG('P.DEUDA')},1,MIN({N},Fin))", fmt=NF_USD)
    nm("r_DeudaFinP", "Control", f"C{r}")
    ws.freeze_panes = "A5"


# ============================================================================ MAIN
def marcar_modificados():
    """Formato condicional: celdas vigentes distintas del presupuesto resaltadas en amarillo/azul."""
    pairs = [("Supuestos", f"E6:E{SUPR['p_last']}", "E6<>D6"),
             ("Productos", f"H{PR['first']}:H{PR['last']}", f"H{PR['first']}<>G{PR['first']}"),
             ("Productos", f"J{PR['first']}:J{PR['last']}", f"J{PR['first']}<>I{PR['first']}"),
             ("Productos", f"L{PR['first']}:L{PR['last']}", f"L{PR['first']}<>K{PR['first']}"),
             ("Productos", f"N{PR['first']}:N{PR['last']}", f"N{PR['first']}<>M{PR['first']}"),
             ("Productos", f"D{PR['vol_V']}:M{PR['vol_V'] + NP - 1}", f"D{PR['vol_V']}<>D{PR['vol_O']}"),
             ("Costos", f"F{CO['per_first']}:F{CO['per_last']}", f"F{CO['per_first']}<>E{CO['per_first']}"),
             ("Costos", f"G{CO['dotv_first']}:P{CO['dotv_first'] + NPER - 1}", f"G{CO['dotv_first']}<>G{CO['per_first']}"),
             ("Costos", f"F{CO['fc_first']}:F{CO['fc_first'] + NFIJ - 1}", f"F{CO['fc_first']}<>E{CO['fc_first']}")]
    for sh, rng, cond in pairs:
        WS[sh].conditional_formatting.add(rng, FormulaRule(formula=[cond], fill=FL_IN, font=Font(name=FONT, size=9, color="0000FF", bold=True)))


def main():
    build_inicio()
    build_supuestos()
    build_productos()
    build_costos()
    build_inversion()
    build_reales()
    build_calculo()
    build_resultados()
    build_control()
    marcar_modificados()
    tabs = {"Inicio": "1F3864", "Supuestos": "FFC000", "Productos": "FFC000", "Costos": "FFC000", "Inversion": "FFC000", "Reales": "FFC000",
            "Calculo": "7F7F7F", "Resultados": "548235", "Control": "C00000"}
    for s in SHEETS:
        WS[s].sheet_properties.tabColor = tabs[s]
        WS[s].sheet_view.zoomScale = 90
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(OUT)
    print("Guardado:", OUT)


if __name__ == "__main__":
    main()
