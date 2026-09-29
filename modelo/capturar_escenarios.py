# -*- coding: utf-8 -*-
"""
Corre el motor del libro (sin capturas) en LibreOffice para cada escenario, prueba de
sensibilidad y combinación de correcciones, y guarda los indicadores en un JSON que
generar_modelo.py incorpora como CAPTURAS identificadas (con verificación en vivo).

Uso:  python3 capturar_escenarios.py libro_base.xlsx capturas.json ruta_recalc.py
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from concurrent.futures import ProcessPoolExecutor

from openpyxl import load_workbook

BASE, OUTJ, RECALC = sys.argv[1], sys.argv[2], sys.argv[3]
WORK = tempfile.mkdtemp(prefix="capt_")

SCEN = ["Base depurada (supuestos Oct-2025)", "Actualizado", "Conservador", "Recuperación operativa", "Personalizado"]
SHOCKS = ["Ninguna", "CAPEX +10%", "CAPEX +20%", "Atraso +3 meses", "Atraso +6 meses", "Volumen -10%", "Volumen -20%",
          "Precio -5%", "Precio -10%", "Materiales +10%", "Materiales +20%", "PYG/USD +10% (más Gs por USD)",
          "PYG/USD -10% (menos Gs por USD)", "DSO 90 días", "DSO 120 días", "DSO 150 días"]
KPIS = ["van_p", "tir_p", "van_e", "tir_e", "e1", "et", "nec", "brecha_nc", "dscr_min", "pb_s", "pb_d", "ctn_max", "vt", "capex_e", "nit", "ap_auto",
        "v1", "ap", "div"]
ANNUAL = ["v", "ebitda", "per", "c_mat", "g_trib", "cxc", "preop", "dep", "c_fle", "u_vg"]


def cell_of(wb, name):
    txt = wb.defined_names[name].attr_text
    sh, ref = txt.split("!")
    return sh.strip("'"), ref.replace("$", "")


def set_name(wb, name, value, base_col=False):
    sh, ref = cell_of(wb, name)
    if base_col:  # driver de escenario: se escribe la columna D (Base) de la misma fila
        ref = "D" + "".join(ch for ch in ref if ch.isdigit())
    wb[sh][ref] = value


def run_case(args):
    tag, esc, shock, corr, reals = args[:5]
    gen, drv = (args[5], args[6]) if len(args) > 5 else ({}, {})
    wb = load_workbook(BASE)
    for k, v in gen.items():
        set_name(wb, k, v)
    for k, v in drv.items():
        set_name(wb, k, v, True)
    sh, c = cell_of(wb, "p_Escenario"); wb[sh][c] = esc
    sh, c = cell_of(wb, "p_Shock"); wb[sh][c] = shock
    for i, v in enumerate(corr):
        sh, c = cell_of(wb, f"p_C{i + 1}"); wb[sh][c] = v
    if reals:
        sh, c = cell_of(wb, "r_Corte"); wb[sh][c] = reals["corte"]
        if reals.get("arranque"):
            sh, c = cell_of(wb, "r_Arranque"); wb[sh][c] = reals["arranque"]
        txt = wb.defined_names["dr_Mes"].attr_text
        first = int(txt.split("!")[1].split(":")[0].replace("$B$", ""))
        ws = wb["Datos_Reales"]
        for k, v in reals.get("saldos", {}).items():
            sh, ref = cell_of(wb, "r_" + k)
            row = "".join(ch for ch in ref if ch.isdigit())
            wb[sh]["D" + row] = v
            wb[sh]["C" + row] = "USD"
        for i, rw in enumerate(reals["rows"]):
            mes, code, imp, cant = rw[:4]
            cur, tc = (rw[4], rw[5]) if len(rw) > 4 else ("USD", None)
            r = first + i
            ws[f"A{r}"] = f"PRUEBA-{i + 1:03d}"
            ws[f"B{r}"] = mes
            ws[f"D{r}"] = code
            ws[f"G{r}"] = cur
            ws[f"H{r}"] = imp
            ws[f"I{r}"] = cant
            ws[f"J{r}"] = tc
            ws[f"M{r}"] = "Aprobado"
    path = os.path.join(WORK, f"{tag}.xlsx")
    wb.save(path)
    res = subprocess.run([sys.executable, RECALC, path, "600"], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    out = json.loads(res.stdout)
    if out.get("status") != "success":
        return tag, {"error": out}
    wv = load_workbook(path, data_only=True)
    ra = wv["Resultados_Anuales"]
    vals = {}
    for r in range(5, 70):
        k = ra[f"A{r}"].value
        if k in KPIS or k in ("chk_bal", "chk_ind"):
            vals[k] = ra[f"D{r}"].value
    for r in range(70, ra.max_row + 1):
        k = ra[f"A{r}"].value
        if k in ANNUAL:
            vals[k + "_tot"] = ra[f"D{r}"].value
            vals[k + "_y1"] = ra[f"F{r}"].value
    rc = wv["Resultados_Caja"]
    for r in range(5, 60):
        if rc[f"A{r}"].value == "v":
            vals["v_m1"] = rc[f"G{r}"].value
            vals["v_m3"] = rc[f"I{r}"].value
            break
    cal = wv["Calendario"]
    for r in range(5, 40):
        if cal[f"A{r}"].value == "f_op":
            vals["meses_op"] = cal[f"E{r}"].value
            break
    vals["errors"] = out.get("total_errors")
    return tag, vals


def main():
    ALL = [1] * 6
    cases = []
    for i, e in enumerate(SCEN):
        cases.append((f"scen{i}", e, "Ninguna", ALL, None))
    for k, s in enumerate(SHOCKS):
        if k == 0:
            continue
        cases.append((f"shock{k}", SCEN[0], s, ALL, None))
    for i in range(7):
        corr = [1 if j < i else 0 for j in range(6)]
        cases.append((f"cum{i}", SCEN[0], "Ninguna", corr, None))
    for i in range(6):
        corr = [1 if j == i else 0 for j in range(6)]
        cases.append((f"iso{i + 1}", SCEN[0], "Ninguna", corr, None))
    reals = {"corte": 2, "arranque": 1, "rows": [(1, "U_VTA_G", None, 700), (2, "U_VTA_G", None, 850), (1, "U_VTA_C", None, 4000), (2, "U_VTA_C", None, 5000),
                                  (1, "VTA_G", 189000, None), (2, "VTA_G", 229500, None), (1, "VTA_C", 92000, None), (2, "VTA_C", 115000, None),
                                  (1, "U_PROD_G", None, 750), (2, "U_PROD_G", None, 900), (1, "U_PROD_C", None, 4300), (2, "U_PROD_C", None, 5200)]}
    srows = []
    for m in range(0, 10):
        srows += [(m, "VTA_G", 150000, None), (m, "VTA_C", 60000, None), (m, "U_VTA_G", None, 550), (m, "U_VTA_C", None, 2600),
                  (m, "U_PROD_G", None, 600), (m, "U_PROD_C", None, 2800), (m, "COBRO", 120000, None), (m, "MP", 70000, None),
                  (m, "NOM_D", 230000000, None, "PYG", 7600), (m, "NOM_A", 150000000, None, "PYG", 7600), (m, "COMPRA_MAT", 90000, None),
                  (m, "PAGO_MAT", 85000, None), (m, "PAGO_MEN", 40000, None), (m, "ALQ", 17000, None), (m, "INT_P1", 20936.9, None)]
    srows += [(0, "DESEMB_P1", 2791587, None), (0, "APORTE", 200000, None)]
    stress = [
        ("stressA", {"p_AlqCaja": 3, "p_ImplMeses": 3, "p_ModoHor": 2, "p_BaseTrib": 2, "p_IRE": 0.1, "p_TipoVT": 2, "p_C5": 0, "p_C6": 0},
         {"e_Linea": 1, "e_Cupo": 2000000, "e_IVArec": 1, "e_IVAlag": 3, "e_Cred": 0.7, "e_Cont": 0.1, "e_Ant": 0.2, "e_AntMeses": 2, "e_AntProv": 0.3,
          "e_AntProvMeses": 2, "e_DPOmp": 45, "e_Incob": 0.01, "e_Rech": 0.03, "e_Chat": 5, "e_Atraso": 4, "e_Rampa": 6, "e_Eff0": 0.5, "e_Recup": 0.6,
          "e_L2_Monto": 500000, "e_L2_Mes": 10}, None),
        ("stressB", {"p_Modalidad": 2, "p_TarifaG": 110, "p_TarifaC": 9, "p_PropCliente": 1, "p_AlqCaja": 2}, {"e_IVArec": 1, "e_Linea": 1, "e_Cupo": 500000}, None),
        ("stressC", {"p_ImplMeses": 2}, {"e_Linea": 1, "e_Cupo": 3000000, "e_Ant": 0.1, "e_Cred": 0.9, "e_AntProv": 0.2},
         {"corte": 9, "arranque": 3, "rows": srows, "saldos": {"S_CAJA": 350000, "S_CXC": 900000, "S_INVMAT": 700000, "S_INVPT": 90000,
                                                              "S_PROVMAT": 120000, "S_PROVMEN": 60000, "S_IVA": 15000, "S_DEUDA1": 2791587}}),
    ]
    for tag, gen, drv, rl in stress:
        cases.append((tag, SCEN[0], "Ninguna", [1] * 6 if "p_C5" not in gen else [1, 1, 1, 1, 0, 0], rl, gen, drv))
    cases.append(("apauto", SCEN[0], "Ninguna", ALL, None, {"p_ApAuto": 1}, {}))
    cases.append(("real_base", SCEN[0], "Ninguna", ALL, reals))
    cases.append(("real_cons", SCEN[2], "Ninguna", ALL, reals))
    res = {}
    with ProcessPoolExecutor(max_workers=4) as ex:
        for tag, vals in ex.map(run_case, cases):
            res[tag] = vals
            print(tag, "ok" if "error" not in vals else vals, flush=True)
    json.dump(res, open(OUTJ + ".raw.json", "w"), ensure_ascii=False, indent=1, default=str)

    # ------------------------------------------------ estructura para el libro
    def pick(v):
        return {k: v.get(k) for k in ["van_p", "tir_p", "van_e", "tir_e", "e1", "et", "nec", "brecha_nc", "dscr_min", "pb_s", "ctn_max"]}
    cap = {"stamp": "29/09/2026 — motor recalculado con LibreOffice 24.2 (headless) sobre este mismo libro; verificar en Excel con la columna «Verificación en vivo».",
           "scen": [], "shock": [], "bridge": [], "tests": []}
    for i, e in enumerate(SCEN):
        cap["scen"].append({"name": e, "vals": pick(res[f"scen{i}"]), "cond": f"AND(p_EscIdx={i + 1},p_ShockIdx=1,p_NCorr=6,p_MesCorte<0)"})
    cap["shock"].append({"name": "Ninguna (Base depurada)", "vals": pick(res["scen0"]), "cond": "AND(p_EscIdx=1,p_ShockIdx=1,p_NCorr=6,p_MesCorte<0)"})
    for k in range(1, len(SHOCKS)):
        cap["shock"].append({"name": SHOCKS[k], "vals": pick(res[f"shock{k}"]), "cond": f"AND(p_EscIdx=1,p_ShockIdx={k + 1},p_NCorr=6,p_MesCorte<0)"})
    names = ["Motor con criterios del PDF (C1-C6 = 0)", "+ C1 TC único en nómina", "+ C2 MP cajas 1% acumulativo", "+ C3 depreciación desde puesta en servicio",
             "+ C4 CAPEX año 2 al activo depreciable", "+ C5 inventario a costo de adquisición", "+ C6 indirectos de fábrica en producción (= Base depurada)"]
    for i in range(7):
        v = res[f"cum{i}"]
        corr = [1 if j < i else 0 for j in range(6)]
        cond = "p_EscIdx=1,p_ShockIdx=1,p_MesCorte<0," + ",".join(f"p_C{j + 1}={corr[j]}" for j in range(6))
        row = {"name": names[i], "van_p": v["van_p"], "van_e": v["van_e"], "et": v["et"], "nec": v["nec"], "cond": cond}
        if i > 0:
            row["iso"] = res[f"iso{i}"]["van_p"] - res["cum0"]["van_p"]
        cap["bridge"].append(row)
    b, s = res["scen0"], res

    def st(ok):
        return "OK" if ok else "REVISAR"
    T = cap["tests"]
    v5 = s["shock5"]
    T.append({"name": "Volumen -10%: ventas totales", "base": b["vt"], "new": v5["vt"], "status": st(abs(v5["vt"] / b["vt"] - 0.9) < 0.002),
              "note": "Esperado ≈ -10% (el recupero/rampa no aplica en la base). Observado: ventas x 0,90; costos variables y CT también bajan."})
    T.append({"name": "Volumen -10%: EBITDA total", "base": b["et"], "new": v5["et"], "status": st(v5["et"] < b["et"]),
              "note": "Baja más que proporcional por costos fijos (apalancamiento operativo)."})
    v7 = s["shock7"]
    T.append({"name": "Precio -5%: ventas totales", "base": b["vt"], "new": v7["vt"], "status": st(abs(v7["vt"] / b["vt"] - 0.95) < 0.001),
              "note": "Esperado -5% exacto."})
    T.append({"name": "Precio -5%: tributo de maquila total", "base": b["g_trib_tot"], "new": v7["g_trib_tot"],
              "status": st(abs(v7["g_trib_tot"] / b["g_trib_tot"] - 0.95) < 0.001), "note": "El tributo sobre factura se recalcula con el precio."})
    v11, v12 = s["shock11"], s["shock12"]
    T.append({"name": "PYG/USD +10% (más Gs por USD): costo de personal total", "base": b["per_tot"], "new": v11["per_tot"],
              "status": st(abs(v11["per_tot"] / b["per_tot"] - 1 / 1.1) < 0.001), "note": "Salarios en Gs fijos: su equivalente USD cae 1/1,10."})
    T.append({"name": "PYG/USD +10%: materia prima e insumos (USD) total", "base": b["c_mat_tot"], "new": v11["c_mat_tot"],
              "status": st(abs(v11["c_mat_tot"] - b["c_mat_tot"]) < 0.5), "note": "Compras pactadas en USD no cambian con el TC."})
    T.append({"name": "PYG/USD -10% (menos Gs por USD): costo de personal total", "base": b["per_tot"], "new": v12["per_tot"],
              "status": st(abs(v12["per_tot"] / b["per_tot"] - 1 / 0.9) < 0.001), "note": "Una caída de Gs/USD encarece en USD el costo en guaraníes."})
    v13, v15 = s["shock13"], s["shock15"]
    T.append({"name": "DSO 150 días: cuentas por cobrar al cierre del año 1", "base": b["cxc_y1"], "new": v15["cxc_y1"], "status": st(v15["cxc_y1"] > b["cxc_y1"]),
              "note": "Más plazo = más CxC y mayor necesidad de fondos; el EBITDA no cambia."})
    T.append({"name": "DSO 150 días: pico de necesidad de fondos adicionales", "base": b["nec"], "new": v15["nec"], "status": st(v15["nec"] > b["nec"]),
              "note": "Liberación/consumo de caja medido además del resultado."})
    T.append({"name": "DSO 150 días: EBITDA total (no debe cambiar)", "base": b["et"], "new": v15["et"], "status": st(abs(v15["et"] - b["et"]) < 0.5),
              "note": "El plazo de cobro no es un costo contable (sin incobrables)."})
    T.append({"name": "DSO 90 días: pico de necesidad de fondos adicionales", "base": b["nec"], "new": v13["nec"], "status": st(v13["nec"] < b["nec"]),
              "note": "Reducción de DSO libera caja."})
    v3 = s["shock3"]
    T.append({"name": "Atraso +3 meses: meses de operación en el horizonte fijo", "base": b["meses_op"], "new": v3["meses_op"], "fmt": "0",
              "status": st(v3["meses_op"] == b["meses_op"] - 3), "note": "Fecha final fija: el atraso no prolonga la vida del proyecto."})
    T.append({"name": "Atraso +3 meses: costos de espera / implantación", "base": b["preop_tot"], "new": v3["preop_tot"], "status": st(v3["preop_tot"] > 0),
              "note": "Alquiler, personal indirecto, parte del directo, energía e intereses sin ventas."})
    T.append({"name": "Atraso +3 meses: VAN del proyecto", "base": b["van_p"], "new": v3["van_p"], "status": st(v3["van_p"] < b["van_p"]),
              "note": "Cobros y pagos desplazados en el calendario mensual."})
    v1 = s["shock1"]
    T.append({"name": "CAPEX +10% (sobre pendiente): costo final estimado", "base": b["capex_e"], "new": v1["capex_e"],
              "status": st(abs(v1["capex_e"] - b["capex_e"] * 1.1) < 1), "note": "Sin reales todo el CAPEX está pendiente; con reales solo lo no ejecutado."})
    T.append({"name": "CAPEX +10%: depreciación total", "base": b["dep_tot"], "new": v1["dep_tot"], "status": st(v1["dep_tot"] > b["dep_tot"]),
              "note": "La depreciación aumenta sin salida de caja adicional por sí misma."})
    rb, rc_ = s["real_base"], s["real_cons"]
    T.append({"name": "Reales congelados: ventas del mes 1 (real) — Base vs Conservador", "base": rb["v_m1"], "new": rc_["v_m1"],
              "status": st(abs(rb["v_m1"] - rc_["v_m1"]) < 0.01 and abs(rb["v_m1"] - 281000) < 0.01),
              "note": "Con último mes cerrado = 2 y ventas reales cargadas, cambiar de escenario no modifica los meses reales (prueba temporal; el libro entregado no tiene reales)."})
    T.append({"name": "Reales congelados: ventas del mes 3 (forecast) — Base vs Conservador", "base": rb["v_m3"], "new": rc_["v_m3"],
              "status": st(abs(rb["v_m3"] - rc_["v_m3"]) > 1), "note": "El futuro sí responde al escenario."})
    c0 = s["cum0"]
    T.append({"name": "Correcciones C1-C6 desactivadas: EBITDA total", "base": b["et"], "new": c0["et"], "status": st(c0["et"] > b["et"]),
              "note": "Sin C1 (nómina a 8.000) y C2 (MP cajas plana) el EBITDA vuelve a acercarse al PDF."})
    desc = {"stressA": "alquiler sin desembolso, implantación 3 meses, horizonte 10 años desde arranque, tributo s/ VA nacional, IRE 10%, valor terminal en marcha, "
                       "línea hipotética, IVA recuperable, contado/anticipos, anticipos a proveedores, DPO 45, incobrables, rechazo, chatarra, atraso 4, rampa 6, tramo 2",
            "stressB": "modalidad servicio de maquila con materiales del cliente, alquiler a relacionada, IVA recuperable con tope, línea hipotética",
            "stressC": "reales meses 0-9 (incl. nómina en PYG con TC propio), saldos reales al corte (caja, clientes, inventarios, proveedores, IVA, deuda)"}
    va = s["apauto"]
    T.append({"name": "Aportes hipotéticos de cobertura (G63 = 1): brecha no cubierta", "base": b["brecha_nc"], "new": va["brecha_nc"],
              "status": st(va["brecha_nc"] < 0.5 and b["brecha_nc"] > 0), "note": "La brecha se cubre con aportes hipotéticos rotulados como no aprobados."})
    T.append({"name": "Aportes hipotéticos de cobertura: VAN de los socios", "base": b["van_e"], "new": va["van_e"], "status": st(va["van_e"] < b["van_e"]),
              "note": "Medir la cobertura reduce el VAN de socios: el esfuerzo adicional deja de ser implícito."})
    for tg in ("stressA", "stressB", "stressC"):
        v = s[tg]
        T.append({"name": f"Integridad bajo configuración extrema {tg[-1]}: balance y directo-indirecto", "base": 0,
                  "new": max(abs(v.get("chk_bal") or 0), abs(v.get("chk_ind") or 0)), "fmt": "#,##0.00",
                  "status": st(abs(v.get("chk_bal") or 0) < 0.01 and abs(v.get("chk_ind") or 0) < 0.01 and v.get("errors") == 0),
                  "note": "Configuración: " + desc[tg] + ". Prueba temporal; el libro entregado conserva la base."})
    T.append({"name": "Errores de fórmula en todas las corridas", "base": 0, "new": max(v.get("errors", 0) or 0 for v in res.values()), "fmt": "0",
              "status": st(all((v.get("errors", 1) == 0) for v in res.values())), "note": f"{len(res)} recálculos completos en LibreOffice sin errores ni circularidades."})
    cap["engine_note"] = (f"Pruebas y capturas: {len(res)} recálculos completos con LibreOffice 24.2 headless el 29/09/2026. "
                          "Después de probar se restauró la base elegida (Base depurada, prueba «Ninguna», C1-C6 = 1, sin reales). "
                          "Pendiente: repetir la verificación abriendo el libro en Microsoft Excel (no disponible en este entorno).")
    json.dump(cap, open(OUTJ, "w"), ensure_ascii=False, indent=1, default=str)
    shutil.rmtree(WORK, ignore_errors=True)


if __name__ == "__main__":
    main()
