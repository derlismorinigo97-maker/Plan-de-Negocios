"""Extracción de texto de archivos cargados (Excel, PDF, Word, CSV, texto) y exportación de resultados."""
import io

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

EXTENSIONES = ["xlsx", "xlsm", "pdf", "docx", "csv", "txt", "md"]


def _excel(datos, max_filas=2000):
    # data_only=True lee los valores calculados (el último guardado en Excel), no las fórmulas.
    wb = load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
    partes = []
    for ws in wb.worksheets:
        filas = []
        for i, fila in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_filas:
                filas.append(f"... (hoja truncada en {max_filas} filas)")
                break
            if any(v is not None for v in fila):
                filas.append("\t".join("" if v is None else str(v) for v in fila))
        if filas:
            partes.append(f"## Hoja: {ws.title}\n" + "\n".join(filas))
    return "\n\n".join(partes)


def extraer_texto(nombre, datos, max_caracteres):
    """Devuelve (texto, aviso). El aviso indica truncamiento o problemas de lectura."""
    ext = nombre.rsplit(".", 1)[-1].lower()
    if ext in ("xlsx", "xlsm"):
        texto = _excel(datos)
    elif ext == "pdf":
        texto = "\n\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(datos)).pages)
    elif ext == "docx":
        texto = "\n".join(p.text for p in Document(io.BytesIO(datos)).paragraphs)
    else:
        texto = datos.decode("utf-8", errors="replace")
    aviso = ""
    if not texto.strip():
        aviso = f"{nombre}: no se pudo extraer texto (¿PDF escaneado o archivo vacío?)."
    elif len(texto) > max_caracteres:
        texto = texto[:max_caracteres]
        aviso = f"{nombre}: truncado a {max_caracteres:,} caracteres para controlar el costo."
    return texto, aviso


def estimar_tokens(texto):
    # Aproximación: ~4 caracteres por token en español. Sirve para estimar costos, no para facturar.
    return len(texto) // 4


def a_docx(titulo, secciones):
    """secciones: lista de (subtítulo, texto markdown). Devuelve bytes del .docx."""
    doc = Document()
    doc.add_heading(titulo, level=1)
    for subtitulo, texto in secciones:
        doc.add_heading(subtitulo, level=2)
        for parrafo in texto.split("\n"):
            doc.add_paragraph(parrafo)
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
